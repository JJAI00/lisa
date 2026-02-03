import tkinter as tk
from tkinter import ttk, messagebox
from tkcalendar import DateEntry
import pandas as pd
import numpy as np
from datetime import datetime, timedelta
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
import os
import re
import hashlib
from core.sendlog import KEY_FILE, SCOPES
from core.data_loader import get_global_auth, safe_api_call

# 제품명에서 줄바꿈 문자 제거 함수 추가
def clean_product_name(name):
    return str(name).replace('\n', ' ').replace('\r', ' ').replace('\u2028', ' ').replace('\u2029', ' ')

class QuoteSimpleDialog:
    """
    견적 작성(기존) 다이얼로그: 과거 견적 내용, 견적 하기(기존) 표 UI 생성
    """
    def __init__(self, parent, past_rows=None, username=None, callback=None):
        print(f"QuoteSimpleDialog 초기화 시작: past_rows={len(past_rows) if past_rows else 0}, username={username}")
        self.parent = parent
        self.past_rows = past_rows or []
        self.username = username or '테스트유저'
        self.callback = callback
        
        # 데이터 초기화
        self.price_df = pd.DataFrame()
        self.categories = []
        self.product_costs = {}
        
        # 행 선택 관련 초기화
        self.selected_row_index = None
        
        # Status 변수 초기화
        self.status_var = tk.StringVar(value="60%")
        
        # 데이터 로드
        print("견적 작성(기존): 가격 데이터 로딩 시작")
        self.load_price_data()
        
        # UI 생성
        print("견적 작성(기존): UI 생성 시작")
        self.win = tk.Toplevel(parent)
        self.win.title("견적 작성 (기존)")
        self.win.geometry("1356x800")  # 창 크기를 더 크게 조정
        self.win.configure(bg="#F7F9FB")
        
        # 모달 설정
        self.win.transient(parent)
        self.win.grab_set()
        
        print("견적 작성(기존): 위젯 생성 시작")
        self.create_widgets()
        print("견적 작성(기존): 창 중앙 배치")
        self.center_window()
        print("견적 작성(기존): 초기화 완료")

    def center_window(self):
        """창을 화면 중앙에 배치"""
        try:
            if not self.win:
                return
                
            # 창이 유효한지 확인
            if not self.win.winfo_exists():
                return
                
            self.win.update_idletasks()
            
            # 창이 파괴되었는지 다시 한번 확인
            if not self.win.winfo_exists():
                return
                
            width = self.win.winfo_width()
            height = self.win.winfo_height()
            
            # 유효한 크기인지 확인
            if width <= 0 or height <= 0:
                return
                
            x = (self.win.winfo_screenwidth() // 2) - (width // 2)
            y = (self.win.winfo_screenheight() // 2) - (height // 2)
            self.win.geometry(f"{width}x{height}+{x}+{y}")
            
        except tk.TclError as e:
            print(f"견적 작성(기존): 창 중앙 배치 중 오류 발생 - {e}")
        except Exception as e:
            print(f"견적 작성(기존): 창 중앙 배치 중 예상치 못한 오류 - {e}")

    def load_price_data(self):
        """price_list에서 대분류, 제품, 원가 정보 로드 (최적화된 방식)"""
        try:
            print("견적 작성(기존): 가격 데이터 로딩 시작...")
            
            # 전역 인증 사용
            gc = get_global_auth()
            
            def load_price():
                print("견적 작성(기존): Google Sheets 연결 중...")
                sh = gc.open_by_key('1OGt__Olrempc5ZUy6QqqhvgRseFLanv1xCPaw5Y2qJk')
                ws = sh.worksheet('Sheet1')
                print("견적 작성(기존): 데이터 읽기 중...")
                # 헤더 중복 문제 해결을 위해 expected_headers 사용
                expected_headers = ['대분류', '소분류', '제품', '상세 제품명', '구분', 'USD', 'EUR', 'DC', '원가', '매입처']
                data = ws.get_all_records(expected_headers=expected_headers)
                print(f"견적 작성(기존): {len(data)}행 데이터 읽기 완료")
                return data
            
            # 직접 API 호출 (타임아웃 로직 제거)
            print("견적 작성(기존): safe_api_call 호출 중...")
            price_data = safe_api_call(load_price)
            
            if price_data is None:
                print("견적 작성(기존): safe_api_call이 None을 반환했습니다.")
                # 직접 호출 시도
                try:
                    print("견적 작성(기존): 직접 API 호출 시도...")
                    price_data = load_price()
                except Exception as direct_error:
                    print(f"견적 작성(기존): 직접 API 호출도 실패: {direct_error}")
                    price_data = None
            
            if price_data:
                print("견적 작성(기존): 데이터프레임 변환 중...")
                self.price_df = pd.DataFrame(price_data)
                self.price_df.columns = self.price_df.columns.str.strip()
                
                # 대분류 목록 (중복 제거)
                self.categories = sorted(self.price_df['대분류'].dropna().unique().tolist())
                # 소분류 목록 (중복 제거)
                self.subcategories = sorted(self.price_df['소분류'].dropna().unique().tolist())
                
                # 제품별 원가 딕셔너리
                self.product_costs = {}
                for _, row in self.price_df.iterrows():
                    product = row.get('제품', '')
                    cost = row.get('원가', 0)
                    if product and cost:
                        self.product_costs[product] = cost
                        
                print(f"견적 작성(기존): 가격 데이터 로드 완료 - {len(self.price_df)}행")
            else:
                print("견적 작성(기존): 가격 데이터를 가져올 수 없어 빈 데이터로 초기화합니다.")
                self.price_df = pd.DataFrame()
                self.categories = []
                self.subcategories = []
                self.product_costs = {}
                    
        except Exception as e:
            print(f"견적 작성(기존): price_list 로드 오류 - {e}")
            print("견적 작성(기존): 빈 데이터로 초기화하여 계속 진행합니다.")
            self.price_df = pd.DataFrame()
            self.categories = []
            self.subcategories = []
            self.product_costs = {}

    def get_subcategories_by_category(self, category):
        """대분류에 따른 소분류 목록 반환"""
        if self.price_df.empty:
            return []
        subcategories = self.price_df[self.price_df['대분류'] == category]['소분류'].dropna().unique().tolist()
        return sorted(subcategories)

    def get_products_by_subcategory(self, subcategory):
        """소분류에 따른 제품 목록 반환"""
        if self.price_df.empty:
            return []
        products = self.price_df[self.price_df['소분류'] == subcategory]['제품'].dropna().unique().tolist()
        return sorted(products)

    def get_products_by_category(self, category):
        """대분류에 따른 제품 목록 반환 (하위 호환성)"""
        if self.price_df.empty:
            return []
        products = self.price_df[self.price_df['대분류'] == category]['제품'].dropna().unique().tolist()
        return sorted(products)

    def create_widgets(self):
        self.quote_rows = []
        # 과거 견적 내용 라벨
        tk.Label(self.win, text="과거 견적 내용", font=("맑은 고딕", 9, "bold")).pack(anchor='w', padx=10, pady=(10,0))
        # 과거 견적 내용 표
        past_frame = tk.Frame(self.win)
        past_frame.pack(fill='x', padx=10)
        self.past_cols = [
            '고객사명','거래처명','계산서 발행일','대분류','제품','수량','판매 단가','판매 합계','만료일','원가','원가계','이익액','이익율'
        ]
        right_align_cols = ['판매 단가','판매 합계','원가','원가계','이익액']
        past_xscroll = tk.Scrollbar(past_frame, orient='horizontal')
        past_yscroll = tk.Scrollbar(past_frame, orient='vertical')
        self.past_tree = ttk.Treeview(
            past_frame, columns=self.past_cols, show='headings', height=5,
            xscrollcommand=past_xscroll.set, yscrollcommand=past_yscroll.set
        )
        past_xscroll.config(command=self.past_tree.xview)
        past_yscroll.config(command=self.past_tree.yview)
        past_xscroll.pack(side='bottom', fill='x')
        past_yscroll.pack(side='right', fill='y')
        for col in self.past_cols:
            anchor = 'e' if col in right_align_cols else 'center'
            self.past_tree.heading(col, text=col)
            self.past_tree.column(col, width=110, anchor=anchor, stretch=False)
        self.past_tree.pack(fill='both', expand=True)
        for row in self.past_rows:
            values = []
            for col in self.past_cols:
                v = row.get(col, "")
                if col in right_align_cols:
                    try:
                        v = f"{int(str(v).replace(',', '')):,}"
                    except:
                        pass
                values.append(v)
            self.past_tree.insert("", "end", values=values)
        # --- 환율 정보와 가격표 환율을 나란히 배치 ---
        self.create_exchange_tables_side_by_side()
        
        # --- 고객사 정보 테이블 ---
        self.create_customer_table()
        
        # --- 견적 하기(기존) 표 ---
        tk.Label(self.win, text="견적 하기(기존)", font=("맑은 고딕", 9, "bold")).pack(anchor='w', padx=10, pady=(20,0))
        # 표를 직접 Frame으로 배치 (Canvas 제거)
        self.table_frame = tk.Frame(self.win)
        self.table_frame.pack(fill='x', padx=10, pady=(0,10))
        self.create_quote_headers()
        for i, past_row in enumerate(self.past_rows):
            self.add_row(past_row)
        btn_frame = tk.Frame(self.win)
        btn_frame.pack(side='bottom', pady=10)
        # 버튼 스타일 통일
        def style_btn(btn, bg, hover_bg):
            btn.configure(font=("맑은 고딕", 10, "bold"), fg="white", bg=bg, relief="flat", padx=12, pady=3)
            btn.bind('<Enter>', lambda e: btn.configure(bg=hover_bg))
            btn.bind('<Leave>', lambda e: btn.configure(bg=bg))
        
        # 행 제어 버튼들 (행 추가 버튼 왼쪽에 배치)
        up_btn = tk.Button(btn_frame, text="⬆️ 선택 행 위로", command=self.move_selected_row_up)
        up_btn.pack(side='left', padx=5)
        style_btn(up_btn, "#3B82F6", "#2563EB")
        
        down_btn = tk.Button(btn_frame, text="⬇️ 선택 행 아래로", command=self.move_selected_row_down)
        down_btn.pack(side='left', padx=5)
        style_btn(down_btn, "#3B82F6", "#2563EB")
        
        delete_btn = tk.Button(btn_frame, text="🗑️ 선택 행 삭제", command=self.delete_selected_row)
        delete_btn.pack(side='left', padx=5)
        style_btn(delete_btn, "#EF4444", "#DC2626")
        
        # 행 추가 버튼
        add_btn = tk.Button(btn_frame, text="행 추가", command=self.add_row)
        add_btn.pack(side='left', padx=5)
        style_btn(add_btn, "#3B82F6", "#2563EB")
        
        # 행 삭제 버튼
        del_btn = tk.Button(btn_frame, text="행 삭제", command=self.remove_row)
        del_btn.pack(side='left', padx=5)
        style_btn(del_btn, "#F59E42", "#EA580C")
        
        # 제품 추가 버튼
        def open_product_add():
            try:
                from quote_product_add import QuoteProductAddWindow
                QuoteProductAddWindow(self.win)
            except Exception as e:
                import traceback
                traceback.print_exc()
                from tkinter import messagebox
                messagebox.showerror('오류', f'제품 추가 창을 여는 데 실패했습니다:\n{e}')
        prod_btn = tk.Button(btn_frame, text="제품 추가", command=open_product_add)
        prod_btn.pack(side='left', padx=5)
        style_btn(prod_btn, "#10B981", "#059669")
        
        # 저장 및 견적서 작성 버튼
        save_btn = tk.Button(btn_frame, text="저장 및 견적서 작성", width=20, command=self.on_save)
        save_btn.pack(side='left', padx=10)
        style_btn(save_btn, "#8B5CF6", "#7C3AED")
        
        # 취소 버튼
        cancel_btn = tk.Button(btn_frame, text="취소", width=10, command=self.win.destroy)
        cancel_btn.pack(side='left', padx=10)
        style_btn(cancel_btn, "#6B7280", "#374151")

    def create_customer_table(self):
        """고객사 추가 테이블 생성 (모던 스타일 적용)"""
        # 고객사 테이블 라벨
        tk.Label(self.win, text="영업현황", font=("맑은 고딕", 11, "bold"), bg='#F5F5F5', fg='#374151').pack(anchor='w', padx=10, pady=(15,5))
        
        # 고객사 테이블 프레임 (카드 스타일)
        customer_frame = tk.Frame(self.win, bg='white', relief='solid', bd=1)
        customer_frame.pack(fill='x', padx=10, pady=(0,10))
        
        # 내부 패딩 프레임
        inner_frame = tk.Frame(customer_frame, bg='white')
        inner_frame.pack(fill='x', padx=15, pady=10)
        
        # 영업사원 행
        salesperson_row = tk.Frame(inner_frame, bg='white')
        salesperson_row.pack(fill='x', pady=3)
        tk.Label(salesperson_row, text="영업사원:", font=("맑은 고딕", 10), width=10, anchor='w', bg='white', fg='#374151').pack(side='left')
        
        # 영업사원 표시 (읽기 전용)
        salesperson_label = tk.Label(salesperson_row, text=self.username, font=("맑은 고딕", 10), fg="#374151", bg="#F9FAFB", relief="solid", bd=1, width=30, anchor='w', padx=5)
        salesperson_label.pack(side='left', padx=(5,0))
        
        # Status 행
        status_row = tk.Frame(inner_frame, bg='white')
        status_row.pack(fill='x', pady=3)
        tk.Label(status_row, text="Status:", font=("맑은 고딕", 10), width=10, anchor='w', bg='white', fg='#374151').pack(side='left')
        
        # Status 선택 드롭다운
        self.status_var = tk.StringVar(value="60%")
        status_combo = ttk.Combobox(status_row, textvariable=self.status_var, width=30, font=("맑은 고딕", 10), state="readonly")
        status_combo['values'] = ['100%', '80%', '60%', '40%', '20%', 'Drop']
        status_combo.pack(side='left', padx=(5,0))
        
        # 예상 발주일 행
        forecast_row = tk.Frame(inner_frame, bg='white')
        forecast_row.pack(fill='x', pady=3)
        tk.Label(forecast_row, text="예상 발주일:", font=("맑은 고딕", 10), width=10, anchor='w', bg='white', fg='#374151').pack(side='left')
        
        # 예상 발주일 입력 필드와 캘린더 버튼
        self.forecast_var = tk.StringVar()
        self.forecast_entry = tk.Entry(forecast_row, textvariable=self.forecast_var, width=30, font=("맑은 고딕", 10), relief='solid', bd=1)
        self.forecast_entry.pack(side='left', padx=(5,0))
        
        # 달력 버튼 (클릭 시 DateEntry 팝업) - 모던 스타일
        calendar_btn = tk.Button(forecast_row, text="📅", command=self.show_calendar, 
                                font=("맑은 고딕", 10), bg="#10B981", fg="white", relief="flat", padx=8, pady=2, cursor="hand2")
        calendar_btn.pack(side='left', padx=(5,0))
        
        # 예상 발주일을 클릭하면 바로 달력이 나오도록 설정
        self.forecast_entry.bind('<Button-1>', lambda e: self.show_calendar())
        
        # 기존 데이터 설정
        if self.past_rows and len(self.past_rows) > 0:
            status_value = self.past_rows[0].get('Status', '')
            forecast_value = self.past_rows[0].get('예상 발주일', '')
            
            if status_value:
                self.status_var.set(status_value)
            
            if forecast_value:
                self.forecast_var.set(forecast_value)

    def create_status_selection_frame(self):
        """Status 선택 프레임 생성"""
        status_frame = tk.LabelFrame(self.win, text="Status 선택", font=("맑은 고딕", 9, "bold"), bg="white", relief="solid", bd=1, labelanchor='nw')
        status_frame.pack(pady=(10, 10), padx=10, fill='x')
        
        # Status 선택 드롭다운
        tk.Label(status_frame, text="Status:", font=("맑은 고딕", 10), fg="#374151", bg="white").pack(side='left', padx=(15, 5), pady=15)
        
        self.status_var = tk.StringVar(value="Drop")
        status_combo = ttk.Combobox(status_frame, textvariable=self.status_var, width=20, font=("맑은 고딕", 10), state="readonly")
        status_combo['values'] = ['100%', '80%', '60%', '40%', '20%', 'Drop']
        status_combo.pack(side='left', padx=(0, 15), pady=15)

    # 헤더 생성 시 self.table_frame에 생성
    def create_quote_headers(self):
        # 표 스타일로 견적 하기(기존) 헤더 생성
        self.col_widths = [10, 10, 50, 3, 10, 13, 13, 10, 13, 10, 5, 10]  # 체크박스 컬럼 제거
        headers = ['대분류', '소분류', '제품', '수량', '판매 단가', '판매 합계', '만료일', '원가', '원가계', '이익액', '이익율', '매입처']
        self.header_labels = []
        
        for i, header in enumerate(headers):
            lbl = tk.Label(self.table_frame, text=header, bg='#F1F5F9', fg="#374151", 
                          relief='solid', font=("맑은 고딕", 10, "bold"), 
                          width=max(self.col_widths[i]//8, 8), height=2, anchor='center')
            lbl.grid(row=0, column=i, sticky='ew')
            self.table_frame.grid_columnconfigure(i, minsize=self.col_widths[i])
            
            # 모든 컬럼에 리사이즈 기능 활성화
            lbl.bind('<Button-1>', lambda e, idx=i: self.start_resize(e, idx))
            lbl.bind('<B1-Motion>', lambda e, idx=i: self.do_resize(e, idx))
            lbl.bind('<ButtonRelease-1>', lambda e, idx=i: self.end_resize(e, idx))
            
            self.header_labels.append(lbl)
        self._resize_col = None
        self._resize_start_x = None

    def start_resize(self, event, col_idx):
        self._resize_col = col_idx
        self._resize_start_x = event.x_root

    def do_resize(self, event, col_idx):
        if self._resize_col is None:
            return
        delta = event.x_root - self._resize_start_x
        # 15px 이상 움직였을 때만 width 1씩 증가/감소, 감도 완화
        if abs(delta) >= 15:
            step = delta // 15
            new_width = max(6, min(40, self.col_widths[col_idx] + step))
            self.col_widths[col_idx] = new_width
            self.header_labels[col_idx].config(width=new_width)
            for row in self.quote_rows:
                widget = self.get_widget_by_col(row, col_idx)
                if widget:
                    widget.config(width=new_width)

    def end_resize(self, event, col_idx):
        self._resize_col = None
        self._resize_start_x = None

    def show_calendar(self, event=None):
        """예상 발주일 선택을 위한 달력 표시"""
        try:
            # 이미 달력 창이 열려있는지 확인
            if hasattr(self, 'calendar_window') and self.calendar_window is not None and self.calendar_window.winfo_exists():
                self.calendar_window.lift()
                self.calendar_window.focus()
                return
            
            # 현재 입력된 날짜가 있으면 파싱 시도
            current_date = self.forecast_var.get().strip()
            if current_date:
                try:
                    import datetime
                    # 다양한 날짜 형식 파싱 시도
                    if '-' in current_date:
                        parsed_date = datetime.datetime.strptime(current_date, '%Y-%m-%d').date()
                    elif '.' in current_date:
                        parsed_date = datetime.datetime.strptime(current_date, '%Y.%m.%d').date()
                    else:
                        parsed_date = datetime.datetime.now().date()
                except:
                    import datetime
                    parsed_date = datetime.datetime.now().date()
            else:
                import datetime
                parsed_date = datetime.datetime.now().date()
            
            # DateEntry 팝업 생성
            self.calendar_window = tk.Toplevel(self.win)
            self.calendar_window.title("예상 발주일 선택")
            self.calendar_window.geometry("400x400")
            self.calendar_window.configure(bg='#F5F5F5')
            self.calendar_window.transient(self.win)
            self.calendar_window.grab_set()
            
            # 창이 닫힐 때 변수 초기화
            def on_closing():
                if self.calendar_window is not None:
                    self.calendar_window.destroy()
                self.calendar_window = None
            
            self.calendar_window.protocol("WM_DELETE_WINDOW", on_closing)
            
            # 달력 위젯 생성 (모던 라이트 테마)
            from tkcalendar import Calendar
            
            # 스타일 설정 - 메인 화면과 통일
            cal = Calendar(self.calendar_window, 
                          selectmode='day',
                          year=parsed_date.year, month=parsed_date.month, day=parsed_date.day,
                          font=("맑은 고딕", 11),
                          # 밝은 테마 색상
                          background='white',
                          foreground='#374151',
                          bordercolor='#E5E7EB',
                          headersbackground='#10B981',  # 청록색 헤더
                          headersforeground='white',
                          selectbackground='#10B981',  # 선택된 날짜 색상
                          selectforeground='white',
                          normalbackground='white',
                          normalforeground='#374151',
                          weekendbackground='#FEF2F2',
                          weekendforeground='#EF4444',
                          othermonthwebackground='#F9FAFB',
                          othermonthweforeground='#9CA3AF')
            cal.pack(pady=20, padx=20, fill="both", expand=True)
            
            # 선택 버튼
            def on_select():
                selected = cal.get_date()
                # yyyy-mm-dd 형식으로 변환
                try:
                    from datetime import datetime
                    # tkcalendar가 반환하는 형식에 따라 처리 (설정에 따라 다를 수 있음)
                    # 보통 'm/d/yy' 형식이 기본값이지만, date_pattern 설정이 없으면 기본값 따름
                    # 안전하게 파싱 시도
                    if '/' in selected:
                        date_obj = datetime.strptime(selected, '%m/%d/%y')
                        formatted = date_obj.strftime('%Y-%m-%d')
                    else:
                        formatted = selected
                except:
                    # 파싱 실패 시 그대로 사용하거나 다른 포맷 시도
                    formatted = selected
                    
                self.forecast_var.set(formatted)
                on_closing()
                
            btn_frame = tk.Frame(self.calendar_window, bg='#F5F5F5')
            btn_frame.pack(pady=10, fill='x')
            
            select_btn = tk.Button(btn_frame, text="선택", command=on_select,
                                  font=("맑은 고딕", 11, "bold"), bg="#10B981", fg="white", 
                                  relief="flat", padx=20, pady=5, cursor="hand2")
            select_btn.pack(side='bottom', pady=10)
            
            # 위치 조정
            try:
                x = self.win.winfo_rootx() + (self.win.winfo_width() // 2) - 200
                y = self.win.winfo_rooty() + (self.win.winfo_height() // 2) - 200
                self.calendar_window.geometry(f"+{x}+{y}")
            except:
                pass
        except Exception as e:
            print(f"달력 표시 오류: {e}")
            import traceback
            traceback.print_exc()

    def get_widget_by_col(self, row_widgets, col_idx):
        col_keys = ['대분류', '소분류', '제품', '수량', '판매 단가', '판매 합계', '만료일', '원가', '원가계', '이익액', '이익율', '매입처']
        key = col_keys[col_idx]
        if key == '대분류':
            return row_widgets.get('대분류_combo')
        if key == '소분류':
            return row_widgets.get('소분류_combo')
        if key == '제품':
            return row_widgets.get('제품_combo')
        return row_widgets.get(f'{key}_entry')

    def add_row(self, past_row=None):
        row_num = len(self.quote_rows) + 1
        row_widgets = {}
        row_widgets['row_index'] = len(self.quote_rows)  # row_index 추가
        
        # 대분류 (드롭다운 + 수동 입력)
        category_var = tk.StringVar()
        category_combo = ttk.Combobox(self.table_frame, textvariable=category_var, 
                                    values=self.categories, width=self.col_widths[0], 
                                    state='normal', font=("맑은 고딕", 10))
        category_combo.grid(row=row_num, column=0, sticky='ew', padx=1, pady=1)
        row_widgets['대분류'] = category_var
        row_widgets['대분류_combo'] = category_combo
        category_combo.bind('<<ComboboxSelected>>', lambda e, row=row_widgets: self.on_category_change(row))
        category_combo.bind('<Return>', lambda e, row=row_widgets: self.on_category_change(row))
        
        # 소분류 (드롭다운 + 수동 입력)
        subcategory_var = tk.StringVar()
        subcategory_combo = ttk.Combobox(self.table_frame, textvariable=subcategory_var, 
                                       width=self.col_widths[1], state='normal', 
                                       font=("맑은 고딕", 10))
        subcategory_combo.grid(row=row_num, column=1, sticky='ew', padx=1, pady=1)
        row_widgets['소분류'] = subcategory_var
        row_widgets['소분류_combo'] = subcategory_combo
        subcategory_combo.bind('<<ComboboxSelected>>', lambda e, row=row_widgets: self.on_subcategory_change(row))
        subcategory_combo.bind('<Return>', lambda e, row=row_widgets: self.on_subcategory_change(row))
        
        # 제품 (드롭다운 + 수동 입력)
        product_var = tk.StringVar()
        product_combo = ttk.Combobox(self.table_frame, textvariable=product_var, 
                                   width=self.col_widths[2], state='normal', 
                                   font=("맑은 고딕", 10))
        product_combo.grid(row=row_num, column=2, sticky='ew', padx=1, pady=1)
        row_widgets['제품'] = product_var
        row_widgets['제품_combo'] = product_combo
        product_combo.bind('<<ComboboxSelected>>', lambda e, row=row_widgets: self.on_product_change(row))
        product_combo.bind('<Return>', lambda e, row=row_widgets: self.on_product_search_and_apply(row))
        
        # 수량 (입력, 가운데 정렬)
        qty_var = tk.StringVar()
        qty_entry = tk.Entry(self.table_frame, textvariable=qty_var, 
                           width=self.col_widths[3], justify='center', 
                           font=("맑은 고딕", 10))
        qty_entry.grid(row=row_num, column=3, sticky='ew', padx=1, pady=1)
        # 수량 입력 시 자동 계산 이벤트 바인딩
        qty_entry.bind('<KeyRelease>', lambda e, row=row_widgets: self.on_qty_change(row))
        qty_entry.bind('<FocusOut>', lambda e, row=row_widgets: self.on_qty_change(row))
        row_widgets['수량'] = qty_var
        row_widgets['수량_entry'] = qty_entry
        
        # 판매 단가 (입력, 오른쪽 정렬)
        price_var = tk.StringVar()
        price_entry = tk.Entry(self.table_frame, textvariable=price_var, 
                             width=self.col_widths[4], justify='right', 
                             font=("맑은 고딕", 10))
        price_entry.grid(row=row_num, column=4, sticky='ew', padx=1, pady=1)
        # 판매 단가 입력 시 자동 계산 이벤트 바인딩
        price_entry.bind('<KeyRelease>', lambda e, row=row_widgets: self.on_price_change(row))
        price_entry.bind('<FocusOut>', lambda e, row=row_widgets: self.on_price_change(row))
        row_widgets['판매 단가'] = price_var
        row_widgets['판매 단가_entry'] = price_entry
        
        # 판매 합계 (자동 계산, 오른쪽 정렬)
        total_var = tk.StringVar()
        total_entry = tk.Entry(self.table_frame, textvariable=total_var, 
                             state='readonly', width=self.col_widths[5], 
                             justify='right', font=("맑은 고딕", 10))
        total_entry.grid(row=row_num, column=5, sticky='ew', padx=1, pady=1)
        row_widgets['판매 합계'] = total_var
        row_widgets['판매 합계_entry'] = total_entry
        
        # 만료일 (입력)
        period_var = tk.StringVar()
        period_value = ''
        if past_row and '만료일' in past_row and past_row['만료일']:
            import datetime
            try:
                expiry = str(past_row['만료일'])
                if isinstance(past_row['만료일'], (datetime.date, datetime.datetime)):
                    expiry_date = past_row['만료일']
                else:
                    expiry_date = datetime.datetime.strptime(expiry, '%Y-%m-%d').date()
                # 기존 만료일에 1년을 더한 새로운 만료일 계산
                new_expiry_date = expiry_date.replace(year=expiry_date.year + 1)
                period_value = new_expiry_date.strftime('%Y-%m-%d')
            except Exception:
                pass
        period_var.set(period_value)
        period_entry = tk.Entry(self.table_frame, textvariable=period_var, 
                              width=self.col_widths[6], font=("맑은 고딕", 10))
        period_entry.grid(row=row_num, column=6, sticky='ew', padx=1, pady=1)
        row_widgets['만료일'] = period_var
        row_widgets['만료일_entry'] = period_entry
        
        # 원가 (입력, 오른쪽 정렬)
        cost_var = tk.StringVar()
        cost_entry = tk.Entry(self.table_frame, textvariable=cost_var, 
                            width=self.col_widths[7], justify='right', 
                            font=("맑은 고딕", 10))
        cost_entry.grid(row=row_num, column=7, sticky='ew', padx=1, pady=1)
        row_widgets['원가'] = cost_var
        row_widgets['원가_entry'] = cost_entry
        
        # 원가계 (자동 계산, 오른쪽 정렬)
        cost_total_var = tk.StringVar()
        cost_total_entry = tk.Entry(self.table_frame, textvariable=cost_total_var, 
                                  state='readonly', width=self.col_widths[8], 
                                  justify='right', font=("맑은 고딕", 10))
        cost_total_entry.grid(row=row_num, column=8, sticky='ew', padx=1, pady=1)
        row_widgets['원가계'] = cost_total_var
        row_widgets['원가계_entry'] = cost_total_entry
        
        # 이익액 (자동 계산, 오른쪽 정렬)
        profit_var = tk.StringVar()
        profit_entry = tk.Entry(self.table_frame, textvariable=profit_var, 
                              state='readonly', width=self.col_widths[9], 
                              justify='right', font=("맑은 고딕", 10))
        profit_entry.grid(row=row_num, column=9, sticky='ew', padx=1, pady=1)
        row_widgets['이익액'] = profit_var
        row_widgets['이익액_entry'] = profit_entry
        
        # 이익율 (자동 계산, 오른쪽 정렬)
        profit_rate_var = tk.StringVar()
        profit_rate_entry = tk.Entry(self.table_frame, textvariable=profit_rate_var, 
                                   state='readonly', width=self.col_widths[10], 
                                   justify='right', font=("맑은 고딕", 10))
        profit_rate_entry.grid(row=row_num, column=10, sticky='ew', padx=1, pady=1)
        row_widgets['이익율'] = profit_rate_var
        row_widgets['이익율_entry'] = profit_rate_entry
        
        # 매입처 (편집 가능, 오른쪽 정렬)
        supplier_var = tk.StringVar()
        supplier_entry = tk.Entry(self.table_frame, textvariable=supplier_var, 
                                width=self.col_widths[11], 
                                justify='right', font=("맑은 고딕", 10))
        supplier_entry.grid(row=row_num, column=11, sticky='ew', padx=1, pady=1)
        row_widgets['매입처'] = supplier_var
        row_widgets['매입처_entry'] = supplier_entry
        
        # 과거 데이터가 있으면 채우기
        if past_row:
            if '대분류' in past_row and past_row['대분류']:
                row_widgets['대분류'].set(past_row['대분류'])
            if '소분류' in past_row and past_row['소분류']:
                row_widgets['소분류'].set(past_row['소분류'])
            if '제품' in past_row and past_row['제품']:
                row_widgets['제품'].set(clean_product_name(past_row['제품']))
            if '수량' in past_row and past_row['수량']:
                row_widgets['수량'].set(str(past_row['수량']))
            if '판매 단가' in past_row and past_row['판매 단가']:
                row_widgets['판매 단가'].set(str(past_row['판매 단가']))
            if '원가' in past_row and past_row['원가']:
                row_widgets['원가'].set(str(past_row['원가']))
            if '매입처' in past_row and past_row['매입처']:
                row_widgets['매입처'].set(str(past_row['매입처']))
                print(f"add_row - past_row에서 매입처 설정: '{past_row['매입처']}'")
        
        self.quote_rows.append(row_widgets)
        
        # 행 선택 바인딩 (행이 추가된 후에 인덱스 설정)
        current_row_index = len(self.quote_rows) - 1
        for key, widget in row_widgets.items():
            if hasattr(widget, 'bind'):
                widget.bind('<Button-1>', lambda e, idx=current_row_index: self.select_row(idx))
        
        self.calculate_row(row_widgets)
        
        # 매입처가 비어 있으면 price_list에서 자동으로 가져오기
        current_supplier = row_widgets['매입처'].get().strip()
        current_product = row_widgets['제품'].get().strip()
        if not current_supplier and current_product:
            product_info = self.get_product_info_from_price_list(current_product)
            if product_info and product_info.get('매입처'):
                row_widgets['매입처'].set(product_info.get('매입처', ''))
                print(f"add_row - price_list에서 매입처 가져옴: '{product_info.get('매입처', '')}'")



    def move_selected_row_up(self):
        """선택된 행을 위로 이동"""
        if (self.selected_row_index is None or 
            self.selected_row_index <= 0 or 
            self.selected_row_index >= len(self.quote_rows)):
            return
        
        target_row = self.selected_row_index - 1
        self.move_row_to_position(self.selected_row_index, target_row)

    def move_selected_row_down(self):
        """선택된 행을 아래로 이동"""
        if (self.selected_row_index is None or 
            self.selected_row_index < 0 or 
            self.selected_row_index >= len(self.quote_rows) - 1):
            return
        
        target_row = self.selected_row_index + 1
        self.move_row_to_position(self.selected_row_index, target_row)

    def move_row_to_position(self, source_row, target_row):
        """행을 특정 위치로 이동"""
        if (source_row == target_row or 
            source_row < 0 or 
            source_row >= len(self.quote_rows) or
            target_row < 0 or 
            target_row >= len(self.quote_rows)):
            return
        
        # 행 데이터 이동
        row_data = self.quote_rows.pop(source_row)
        self.quote_rows.insert(target_row, row_data)
        
        # 모든 행의 row_index 업데이트
        for i, row in enumerate(self.quote_rows):
            row['row_index'] = i
        
        # 위젯 위치 재배치 및 이벤트 바인딩 재설정
        self.rearrange_widgets()
        
        # 선택된 행 인덱스 업데이트
        if self.selected_row_index is not None:
            if self.selected_row_index == source_row:
                self.selected_row_index = target_row
            elif source_row < self.selected_row_index <= target_row:
                self.selected_row_index -= 1
            elif target_row <= self.selected_row_index < source_row:
                self.selected_row_index += 1
            
            # 선택된 행 하이라이트 유지
            if 0 <= self.selected_row_index < len(self.quote_rows):
                self.highlight_row_selection(self.selected_row_index)

    def on_category_change(self, row_widgets):
        """대분류 변경 시 해당 행의 소분류와 제품 콤보박스 업데이트"""
        category = row_widgets['대분류'].get()
        
        # 소분류 목록 업데이트
        subcategories = self.get_subcategories_by_category(category)
        subcategory_combo = row_widgets.get('소분류_combo')
        if subcategory_combo:
            subcategory_combo['values'] = subcategories
            subcategory_combo.set('')  # 소분류 초기화
        
        # 제품 목록 업데이트 (대분류에 따른 전체 제품)
        products = self.get_products_by_category(category)
        product_combo = row_widgets.get('제품_combo')
        if product_combo:
            product_combo['values'] = products
            # 현재 제품이 새로운 대분류에 있으면 유지, 없으면 초기화
            current_product = row_widgets['제품'].get()
            if current_product in products:
                product_combo.set(current_product)
                # 현재 제품의 원가 다시 설정
                self.on_product_change(row_widgets)
            else:
                product_combo.set('')  # 선택 초기화
                row_widgets['원가'].set('')
        else:
            # 원가 초기화
            row_widgets['원가'].set('')
        self.calculate_row(row_widgets)

    def on_subcategory_change(self, row_widgets):
        """소분류 변경 시 해당 행의 제품 콤보박스 업데이트"""
        subcategory = row_widgets['소분류'].get()
        category = row_widgets['대분류'].get()
        
        if subcategory:
            # 소분류에 따른 제품 목록 업데이트
            products = self.get_products_by_subcategory(subcategory)
        else:
            # 소분류가 비어있으면 대분류에 따른 전체 제품 목록
            products = self.get_products_by_category(category)
        
        product_combo = row_widgets.get('제품_combo')
        if product_combo:
            product_combo['values'] = products
            # 현재 제품이 새로운 소분류에 있으면 유지, 없으면 초기화
            current_product = row_widgets['제품'].get()
            if current_product in products:
                product_combo.set(current_product)
                # 현재 제품의 원가 다시 설정
                self.on_product_change(row_widgets)
            else:
                product_combo.set('')  # 선택 초기화
                row_widgets['원가'].set('')
        else:
            # 원가 초기화
            row_widgets['원가'].set('')
        self.calculate_row(row_widgets)

    def get_valid_price(self, val):
        print(f"get_valid_price: Input val = '{val}' (type: {type(val)})")
        try:
            if val is None:
                print("get_valid_price: val is None, returning ''")
                return ''
            if isinstance(val, float) and pd.isna(val):
                print("get_valid_price: val is NaN, returning ''")
                return ''
            val_str = str(val).strip()
            print(f"get_valid_price: val_str after strip = '{val_str}'")
            if val_str == '' or val_str.lower() == 'nan':
                print("get_valid_price: val_str is empty or 'nan', returning ''")
                return ''
            # 통화 기호, 공백 등 제거 (숫자, 소수점만 남김)
            val_str = re.sub(r'[^\d.]', '', val_str)
            print(f"get_valid_price: val_str after regex = '{val_str}'")
            if val_str == '':
                print("get_valid_price: val_str after regex is empty, returning ''")
                return ''
            print(f"get_valid_price: Returning '{val_str}'")
            return val_str
        except Exception as e:
            print(f"get_valid_price: Error during processing '{val}' - {e}, returning ''")
            return ''

    def on_product_change(self, row_widgets, preserve_existing=False):
        """제품 변경 시 원가를 price_list의 '원가' 열에서 가져옴
        
        Args:
            row_widgets: 행 위젯들
            preserve_existing: True이면 기존 원가/매입처 값이 있을 경우 유지
        """
        product = row_widgets['제품'].get()
        category = row_widgets['대분류'].get()
        
        # 기존 원가와 매입처 값 저장
        existing_cost = row_widgets['원가'].get().strip()
        existing_supplier = row_widgets['매입처'].get().strip()
        
        cost = ''
        
        print(f"견적 작성(기존): 제품 변경 - '{product}', 대분류: '{category}'")
        
        if not self.price_df.empty and product:
            # 1. 먼저 현재 대분류에서 제품 검색
            row = self.price_df[(self.price_df['대분류'] == category) & (self.price_df['제품'] == product)]
            
            if not row.empty:
                # 현재 대분류에서 제품을 찾은 경우
                raw_cost = row.iloc[0].get('원가', '')
                cost = self.get_valid_price(raw_cost)
                print(f"견적 작성(기존): 현재 대분류에서 원가 찾음 - 원본값: '{raw_cost}', 변환값: '{cost}'")
                
                # 매입처 정보 설정
                supplier = row.iloc[0].get('매입처', '')
                if supplier:
                    row_widgets['매입처'].set(supplier)
                    print(f"견적 작성(기존): 현재 대분류에서 매입처 설정 - '{supplier}'")
            else:
                # 2. 전체 제품에서 정확히 일치하는 제품 검색
                exact_match = self.price_df[self.price_df['제품'] == product]
                if not exact_match.empty:
                    matched_row = exact_match.iloc[0]
                    matched_category = matched_row['대분류']
                    raw_cost = matched_row.get('원가', '')
                    cost = self.get_valid_price(raw_cost)
                    
                    # 대분류 자동 설정
                    category_combo = row_widgets.get('대분류_combo')
                    if category_combo:
                        category_combo.set(matched_category)
                        row_widgets['대분류'].set(matched_category)
                    
                    # 소분류 자동 설정
                    subcategory = matched_row.get('소분류', '')
                    subcategory_combo = row_widgets.get('소분류_combo')
                    if subcategory_combo and subcategory:
                        subcategory_combo.set(subcategory)
                        row_widgets['소분류'].set(subcategory)
                    
                    # 매입처 자동 설정
                    supplier = matched_row.get('매입처', '')
                    if supplier:
                        row_widgets['매입처'].set(supplier)
                        print(f"견적 작성(기존): 전체 제품에서 매입처 설정 - '{supplier}'")
                    
                    print(f"견적 작성(기존): 전체 제품에서 원가 찾음 - 원본값: '{raw_cost}', 변환값: '{cost}', 대분류: {matched_category}")
                else:
                    # 3. 제품명에 포함되는 제품 검색
                    matching_rows = self.price_df[self.price_df['제품'].str.contains(product, case=False, na=False, regex=False)]
                    if not matching_rows.empty:
                        matched_row = matching_rows.iloc[0]
                        matched_category = matched_row['대분류']
                        matched_product = matched_row['제품']
                        raw_cost = matched_row.get('원가', '')
                        cost = self.get_valid_price(raw_cost)
                        
                        # 대분류 자동 설정
                        category_combo = row_widgets.get('대분류_combo')
                        if category_combo:
                            category_combo.set(matched_category)
                            row_widgets['대분류'].set(matched_category)
                        
                        # 소분류 자동 설정
                        subcategory = matched_row.get('소분류', '')
                        subcategory_combo = row_widgets.get('소분류_combo')
                        if subcategory_combo and subcategory:
                            subcategory_combo.set(subcategory)
                            row_widgets['소분류'].set(subcategory)
                        
                        # 제품명 정확히 설정
                        product_combo = row_widgets.get('제품_combo')
                        if product_combo:
                            product_combo.set(matched_product)
                            row_widgets['제품'].set(matched_product)
                        
                        # 매입처 자동 설정
                        supplier = matched_row.get('매입처', '')
                        if supplier:
                            row_widgets['매입처'].set(supplier)
                            print(f"견적 작성(기존): 부분 일치로 매입처 설정 - '{supplier}'")
                        
                        print(f"견적 작성(기존): 부분 일치로 원가 찾음 - 원본값: '{raw_cost}', 변환값: '{cost}', 제품: {matched_product}")
                    else:
                        print(f"견적 작성(기존): 제품 '{product}'에 대한 원가 정보를 찾을 수 없습니다.")
        else:
            if self.price_df.empty:
                print("견적 작성(기존): 가격 데이터가 비어있습니다.")
            if not product:
                print("견적 작성(기존): 제품명이 비어있습니다.")
        
        # 원가를 항상 천 단위 콤마로 표시
        try:
            if cost and cost != '':
                cost_val = float(cost)
                cost = f"{int(cost_val):,}"
                print(f"견적 작성(기존): 원가 천 단위 콤마 변환: {cost}")
        except Exception as e:
            print(f"견적 작성(기존): 원가 천 단위 콤마 변환 오류: {e}")
            cost = ''
        
        # preserve_existing이 True이고 기존 값이 있으면 유지
        if preserve_existing and existing_cost:
            print(f"견적 작성(기존): 기존 원가 유지: '{existing_cost}'")
        else:
            print(f"견적 작성(기존): 최종 원가 설정: '{cost}'")
            row_widgets['원가'].set(str(cost))
        
        # 매입처 정보 자동 설정 (기존 값이 없거나 preserve_existing이 False인 경우에만)
        if product:
            # preserve_existing이 True이고 기존 매입처가 있으면 유지
            if preserve_existing and existing_supplier:
                print(f"견적 작성(기존): 기존 매입처 유지 - '{existing_supplier}'")
            else:
                # 매입처가 아직 설정되지 않은 경우에만 설정
                current_supplier = row_widgets['매입처'].get().strip()
                if not current_supplier:
                    product_info = self.get_product_info_from_price_list(product)
                    if product_info and product_info.get('매입처'):
                        supplier = product_info.get('매입처', '')
                        row_widgets['매입처'].set(supplier)
                        print(f"견적 작성(기존): 매입처 자동 설정 - '{supplier}'")
                    else:
                        print(f"견적 작성(기존): 제품 '{product}'에 대한 매입처 정보를 찾을 수 없습니다.")
        
        self.calculate_row(row_widgets)

    def on_price_change(self, row_widgets):
        """판매 단가 변경 시 자동 계산"""
        try:
            # 수량과 판매 단가 가져오기
            qty = int(row_widgets['수량'].get().replace(',', '') or 0)
            price_str = row_widgets['판매 단가'].get().replace(',', '')
            price = float(price_str) if price_str else 0
            
            # 판매 합계 계산
            total = qty * price
            row_widgets['판매 합계'].set(f"{int(total):,}" if total > 0 else "")
            
            # 원가계와 이익액도 재계산
            cost = float(row_widgets['원가'].get().replace(',', '') or 0)
            cost_total = qty * cost
            row_widgets['원가계'].set(f"{int(cost_total):,}" if cost_total > 0 else "")
            
            # 이익액 계산
            profit = total - cost_total
            row_widgets['이익액'].set(f"{int(profit):,}" if profit != 0 else "")
            
            # 이익율 계산
            if total > 0:
                profit_rate = (profit / total) * 100
                row_widgets['이익율'].set(f"{profit_rate:.2f}%" if profit_rate != 0 else "")
            else:
                row_widgets['이익율'].set("")
                
        except Exception as e:
            print(f"판매 단가 자동 계산 오류: {e}")
            pass

    def on_qty_change(self, row_widgets):
        """수량 변경 시 자동 계산"""
        try:
            # 수량과 판매 단가 가져오기
            qty_str = row_widgets['수량'].get().replace(',', '')
            qty = int(qty_str) if qty_str else 0
            price = float(row_widgets['판매 단가'].get().replace(',', '') or 0)
            
            # 판매 합계 계산
            total = qty * price
            row_widgets['판매 합계'].set(f"{int(total):,}" if total > 0 else "")
            
            # 원가계와 이익액도 재계산
            cost = float(row_widgets['원가'].get().replace(',', '') or 0)
            cost_total = qty * cost
            row_widgets['원가계'].set(f"{int(cost_total):,}" if cost_total > 0 else "")
            
            # 이익액 계산
            profit = total - cost_total
            row_widgets['이익액'].set(f"{int(profit):,}" if profit != 0 else "")
            
            # 이익율 계산
            if total > 0:
                profit_rate = (profit / total) * 100
                row_widgets['이익율'].set(f"{profit_rate:.2f}%" if profit_rate != 0 else "")
            else:
                row_widgets['이익율'].set("")
                
        except Exception as e:
            print(f"수량 자동 계산 오류: {e}")
            pass

    def calculate_row(self, row_widgets):
        """행의 자동 계산"""
        try:
            # 수량과 판매 단가 가져오기
            qty = float(row_widgets['수량'].get().replace(',', '') or 0)
            price = float(row_widgets['판매 단가'].get().replace(',', '') or 0)
            cost = float(row_widgets['원가'].get().replace(',', '') or 0)
            # 판매 단가 천 단위 콤마 표시
            if price > 0:
                row_widgets['판매 단가'].set(f"{int(price):,}")
            # 판매 합계 계산
            total = qty * price
            row_widgets['판매 합계'].set(f"{int(total):,}" if total > 0 else "")
            # 원가계 계산
            cost_total = qty * cost
            row_widgets['원가계'].set(f"{int(cost_total):,}" if cost_total > 0 else "")
            # 이익액 계산
            profit = total - cost_total
            row_widgets['이익액'].set(f"{int(profit):,}" if profit != 0 else "")
            # 이익율 계산
            if total > 0:
                profit_rate = (profit / total) * 100
                row_widgets['이익율'].set(f"{profit_rate:.2f}%" if profit_rate != 0 else "")
            else:
                row_widgets['이익율'].set("")
        except ValueError:
            pass

    def on_product_search_and_apply(self, row_widgets):
        """엔터 키로 제품 검색 및 적용"""
        if not self.price_df.empty:
            product_input = row_widgets['제품'].get().strip()
            if len(product_input) >= 2:  # 2글자 이상 입력 시 검색 시작
                # 제품명으로 검색 (대분류 무관)
                matching_rows = self.price_df[self.price_df['제품'].str.contains(product_input, case=False, na=False, regex=False)]
                if not matching_rows.empty:
                    # 검색 결과를 콤보박스에 표시
                    product_combo = row_widgets.get('제품_combo')
                    if product_combo:
                        # 검색된 제품명들을 콤보박스 값으로 설정
                        matching_products = matching_rows['제품'].unique().tolist()
                        product_combo['values'] = matching_products
                        
                        # 첫 번째 매칭되는 제품을 미리보기로 표시
                        if matching_products:
                            product_combo.set(matching_products[0])
                            # 드롭다운 리스트 표시
                            product_combo.event_generate('<Down>')
                            product_combo.event_generate('<Up>')

    def remove_row(self):
        """마지막 행 삭제 (기존 기능 유지)"""
        if len(self.quote_rows) > 1:  # 최소 1행은 유지
            # 마지막 행의 위젯들 제거
            last_row = len(self.quote_rows)
            for widget in self.table_frame.winfo_children():
                if widget.grid_info()['row'] == last_row:
                    widget.destroy()
            self.quote_rows.pop()
            
            # 선택된 행이 삭제된 행이었다면 선택 상태 초기화
            if self.selected_row_index == len(self.quote_rows):
                self.selected_row_index = None

    def on_save(self):
        import gspread
        from google.oauth2.service_account import Credentials
        from datetime import datetime
        import openpyxl
        import os
        import shutil
        import subprocess
        from tkinter import messagebox

        # 견적서 ID 생성
        def generate_quote_id():
            """견적서 고유 ID 생성 (날짜_순번 형태)"""
            today = datetime.today()
            date_str = today.strftime('%Y-%m-%d')
            
            # 견적서 번호 생성 (오늘 날짜의 전체 견적서 순번)
            quote_number = self._get_next_quote_number(date_str)
            
            quote_id = f"{date_str}_{quote_number:03d}"
            return quote_id, quote_number

        # 여러 행 선택 시 고객사/거래처 검증
        if len(self.quote_rows) > 1:
            first_customer = self.past_rows[0].get('고객사명', '') if self.past_rows else ''
            first_client = self.past_rows[0].get('거래처명', '') if self.past_rows else ''
            
            for i, past_row in enumerate(self.past_rows[1:], 1):
                current_customer = past_row.get('고객사명', '')
                current_client = past_row.get('거래처명', '')
                
                if current_customer != first_customer or current_client != first_client:
                    messagebox.showwarning(
                        "선택 오류", 
                        "같은 고객사와 거래처의 건을 선택해 주세요.\n\n"
                        f"첫 번째 행: {first_customer} - {first_client}\n"
                        f"{i+1}번째 행: {current_customer} - {current_client}"
                    )
                    return

        # 입력값 검증
        # Status 검증
        selected_status = self.status_var.get()
        if selected_status == "Drop":
            messagebox.showwarning("입력 오류", "영업 활동 Status를 기입해 주세요.")
            return

        # 견적 행 데이터 검증
        if not self.quote_rows:
            messagebox.showwarning("입력 오류", "최소 하나의 제품을 추가해주세요.")
            return

        # 각 견적 행 검증
        for i, row_widgets in enumerate(self.quote_rows):
            unit_price = row_widgets['판매 단가'].get().strip()
            cost = row_widgets['원가'].get().strip()
            
            if not unit_price or float(unit_price.replace(',', '') or 0) <= 0:
                messagebox.showwarning("입력 오류", f"{i+1}번째 행의 판매 단가를 입력해주세요.")
                return
            
            if not cost or float(cost.replace(',', '') or 0) <= 0:
                messagebox.showwarning("입력 오류", f"{i+1}번째 행의 원가를 입력해주세요.")
                return

        # 견적 작성(기존)은 renewal_list에서 정보를 가져와서 새로운 견적서를 만드는 것이므로
        # 항상 새로운 견적서ID를 생성해야 함
        quote_id, quote_number = generate_quote_id()
        print(f"견적 작성(기존) - 새 견적서 ID 생성: {quote_id} (번호: {quote_number})")

        # 1. 엑셀 파일 저장 (quote_origin_01.xlsx → 고객사명_거래처명_대분류_일자.xlsx)
        try:
            # 파일명 생성
            today = datetime.today().strftime('%Y%m%d')
            cust = self.past_rows[0].get('고객사명','') if self.past_rows else ''
            client = self.past_rows[0].get('거래처명','') if self.past_rows else ''
            main_category = self.quote_rows[0]['대분류'].get() if self.quote_rows else ''
            if cust == client or not client:
                filename = f"{cust}_{main_category}_{today}.xlsx"
            else:
                filename = f"{cust}-{client}_{main_category}_{today}.xlsx"
            out_dir = os.path.join(os.getcwd(), 'quote')
            os.makedirs(out_dir, exist_ok=True)
            out_path = os.path.join(out_dir, filename)
            template_path = os.path.join('quote', 'quote_origin_01.xlsx')
            shutil.copy(template_path, out_path)
            wb = openpyxl.load_workbook(out_path)
            ws = wb.active
            # C5: 고객사명 또는 고객사명-거래처명 (병합 셀 대응)
            # openpyxl의 병합 셀 구조를 고려해 top-left 셀에만 값 입력
            c5_row, c5_col = 5, 3
            c5_coord = ws.cell(row=c5_row, column=c5_col).coordinate
            target_value = cust if cust == client or not client else f"{cust} - {client}"
            merged = False
            for merged_range in ws.merged_cells.ranges:
                if c5_coord in merged_range:
                    tl_cell = merged_range.min_row, merged_range.min_col
                    ws.cell(row=tl_cell[0], column=tl_cell[1], value=target_value)
                    merged = True
                    break
            if not merged:
                ws.cell(row=c5_row, column=c5_col, value=target_value)
            # D14, L14, M14, Q14, T14부터 행별로 입력
            start_row = 14
            for i, row in enumerate(self.quote_rows):
                d_row = start_row + i
                # 제품명은 문자열 그대로 (줄바꿈 문자 제거)
                ws[f'D{d_row}'] = clean_product_name(row['제품'].get())
                # 수량, 판매 단가, 원가는 숫자형으로 변환 후 입력
                def parse_number(val):
                    try:
                        return int(str(val).replace(',', ''))
                    except:
                        try:
                            return float(str(val).replace(',', ''))
                        except:
                            return val
                ws[f'L{d_row}'] = parse_number(row['수량'].get())
                ws[f'M{d_row}'] = parse_number(row['판매 단가'].get())
                ws[f'T{d_row}'] = parse_number(row['원가'].get())
                period = row['만료일'].get()
                if period:
                    ws[f'Q{d_row}'] = period
            # 워크시트 이름에서 작은따옴표 제거
            for sheet in wb.sheetnames:
                if sheet.startswith("'"):
                    # 작은따옴표로 시작하는 워크시트 이름 수정
                    new_name = sheet.lstrip("'")
                    wb[sheet].title = new_name
            
            wb.save(out_path)
            # 엑셀 열기 (닫지 않음)
            try:
                os.startfile(out_path)
            except AttributeError:
                subprocess.Popen(['open', out_path])
        except Exception as e:
            import traceback
            traceback.print_exc()
            messagebox.showerror('엑셀 저장 오류', str(e))

        # 2. quote_list 워크시트(구글시트)에 저장 (기존 로직 유지)
        try:
            creds = Credentials.from_service_account_file(KEY_FILE, scopes=SCOPES)
            gc = gspread.authorize(creds)
            sh = gc.open_by_key('1xjIjFe1Q9dq2zeQOkBSpbqIUOyXAxbyTO8lKN-r1lw0')
            ws = sh.worksheet('Sheet1')
            # 헤더 정규화(앞뒤 공백 제거 및 중간 공백 제거 버전도 준비)
            raw_header = ws.row_values(1)
            sheet_header = [h.strip() for h in raw_header]
            print(f"quote_list 헤더: {sheet_header}")

            def normalize(name: str) -> str:
                try:
                    return str(name).replace(' ', '').strip()
                except Exception:
                    return str(name)

            normalized_map = {normalize(h): idx for idx, h in enumerate(sheet_header)}

            expected_fields = ['ID','견적서ID','일자','영업사원','FCST','예상 발주일','Status','고객사명','거래처명','대분류','소분류','제품','수량','판매 단가','판매 합계','원가','원가계','이익액','이익율','만료일','구분','USD','EUR','DC','매입처']
            col_idx = {k: normalized_map.get(normalize(k)) for k in expected_fields}
            print(f"컬럼 인덱스 매핑: {col_idx}")
            cust = self.past_rows[0].get('고객사명','') if self.past_rows else ''
            client = self.past_rows[0].get('거래처명','') if self.past_rows else ''
            today = datetime.today().strftime('%Y-%m-%d')
            user = self.username  # 이미 __init__에서 기본값 설정됨
            # 현재 시트의 마지막 데이터 행 찾기
            last_row = len(ws.get_all_values())
            for i, row in enumerate(self.quote_rows):
                # 필수값 체크: 대분류, 제품, 수량, 판매 단가, 원가
                if not (row['대분류'].get() and row['제품'].get() and row['수량'].get() and row['판매 단가'].get() and row['원가'].get()):
                    continue
                data = [''] * len(sheet_header)
                
                # 견적서 ID (첫 번째 열, 모든 행에 같은 ID)
                if col_idx.get('견적서ID') is not None:
                    data[col_idx['견적서ID']] = quote_id
                    print(f"견적서 ID 저장: {quote_id} -> 행 {i+1}")
                elif col_idx.get('ID') is not None:
                    data[col_idx['ID']] = quote_id
                    print(f"ID 저장: {quote_id} -> 행 {i+1}")
                
                if col_idx['일자'] is not None:
                    data[col_idx['일자']] = today
                if col_idx['영업사원'] is not None:
                    data[col_idx['영업사원']] = user
                if col_idx['FCST'] is not None:
                    data[col_idx['FCST']] = self.forecast_var.get()
                elif col_idx.get('예상 발주일') is not None:
                    data[col_idx['예상 발주일']] = self.forecast_var.get()
                if col_idx['대분류'] is not None:
                    data[col_idx['대분류']] = row['대분류'].get()
                if col_idx['소분류'] is not None:
                    subcategory_value = row['소분류'].get()
                    data[col_idx['소분류']] = subcategory_value
                    print(f"소분류 저장: '{subcategory_value}' -> 인덱스 {col_idx['소분류']}")
                else:
                    print(f"소분류 컬럼을 찾을 수 없음 - col_idx: {col_idx}")
                if col_idx['제품'] is not None:
                    data[col_idx['제품']] = clean_product_name(row['제품'].get())
                if col_idx['수량'] is not None:
                    data[col_idx['수량']] = row['수량'].get()
                if col_idx['판매 단가'] is not None:
                    data[col_idx['판매 단가']] = row['판매 단가'].get()
                if col_idx['원가'] is not None:
                    data[col_idx['원가']] = row['원가'].get()
                if col_idx['고객사명'] is not None:
                    data[col_idx['고객사명']] = cust
                if col_idx['거래처명'] is not None:
                    data[col_idx['거래처명']] = client
                if col_idx['만료일'] is not None:
                    data[col_idx['만료일']] = row['만료일'].get()
                
                # renewal_list에서 매입처 정보 우선 가져오기
                renewal_supplier = None
                if self.past_rows and len(self.past_rows) > i:
                    renewal_supplier = self.past_rows[i].get('매입처', '')
                    print(f"renewal_list에서 매입처 정보: {renewal_supplier}")
                
                # price_list에서 제품 정보 가져와서 저장
                product_name = clean_product_name(row['제품'].get())
                if product_name:
                    product_info = self.get_product_info_from_price_list(product_name)
                    if product_info:
                        print(f"제품 정보 저장 시작: {product_name}")
                        print(f"찾은 제품 정보: {product_info}")
                        print(f"컬럼 인덱스: {col_idx}")
                        
                        # 안전하게 컬럼 존재 여부 확인 후 저장
                        if '구분' in col_idx and col_idx['구분'] is not None:
                            data[col_idx['구분']] = product_info.get('구분', '')
                            print(f"구분 저장: {product_info.get('구분', '')} -> 인덱스 {col_idx['구분']}")
                        if 'USD' in col_idx and col_idx['USD'] is not None:
                            data[col_idx['USD']] = product_info.get('USD', '')
                            print(f"USD 저장: {product_info.get('USD', '')} -> 인덱스 {col_idx['USD']}")
                        if 'EUR' in col_idx and col_idx['EUR'] is not None:
                            data[col_idx['EUR']] = product_info.get('EUR', '')
                            print(f"EUR 저장: {product_info.get('EUR', '')} -> 인덱스 {col_idx['EUR']}")
                        if 'DC' in col_idx and col_idx['DC'] is not None:
                            data[col_idx['DC']] = product_info.get('DC', '')
                            print(f"DC 저장: {product_info.get('DC', '')} -> 인덱스 {col_idx['DC']}")
                        
                        # 매입처: UI 입력값 우선, 없으면 renewal_list, 마지막으로 price_list 사용
                        if '매입처' in col_idx and col_idx['매입처'] is not None:
                            ui_supplier = row['매입처'].get().strip()
                            if ui_supplier:
                                data[col_idx['매입처']] = ui_supplier
                                print(f"매입처 저장 (UI): {ui_supplier} -> 인덱스 {col_idx['매입처']}")
                            elif renewal_supplier:
                                data[col_idx['매입처']] = renewal_supplier
                                print(f"매입처 저장 (renewal_list): {renewal_supplier} -> 인덱스 {col_idx['매입처']}")
                            else:
                                data[col_idx['매입처']] = product_info.get('매입처', '')
                                print(f"매입처 저장 (price_list): {product_info.get('매입처', '')} -> 인덱스 {col_idx['매입처']}")
                        
                        print(f"최종 저장될 데이터: {data}")
                    else:
                        print(f"제품 정보를 찾지 못함: {product_name}")
                        # 제품 정보가 없어도 UI 입력값이나 renewal_list의 매입처 정보는 저장
                        if '매입처' in col_idx and col_idx['매입처'] is not None:
                            ui_supplier = row['매입처'].get().strip()
                            if ui_supplier:
                                data[col_idx['매입처']] = ui_supplier
                                print(f"매입처 저장 (UI): {ui_supplier} -> 인덱스 {col_idx['매입처']}")
                            elif renewal_supplier:
                                data[col_idx['매입처']] = renewal_supplier
                                print(f"매입처 저장 (renewal_list만): {renewal_supplier} -> 인덱스 {col_idx['매입처']}")
                else:
                    print(f"제품명이 비어있음")
                    # 제품명이 없어도 UI 입력값이나 renewal_list의 매입처 정보는 저장
                    if '매입처' in col_idx and col_idx['매입처'] is not None:
                        ui_supplier = row['매입처'].get().strip()
                        if ui_supplier:
                            data[col_idx['매입처']] = ui_supplier
                            print(f"매입처 저장 (UI): {ui_supplier} -> 인덱스 {col_idx['매입처']}")
                        elif renewal_supplier:
                            data[col_idx['매입처']] = renewal_supplier
                            print(f"매입처 저장 (renewal_list만): {renewal_supplier} -> 인덱스 {col_idx['매입처']}")
                
                # Status 컬럼 추가
                if 'Status' in col_idx and col_idx['Status'] is not None:
                    selected_status = self.status_var.get()
                    # Status가 비어있으면 기본값 60% 사용, Drop이면 Drop 그대로 사용
                    if not selected_status:
                        selected_status = "60%"
                        print(f"Status가 비어있으므로 기본값 60% 사용")
                    data[col_idx['Status']] = selected_status
                    print(f"Status 저장: {selected_status} -> 인덱스 {col_idx['Status']}")
                else:
                    print(f"Status 컬럼을 찾을 수 없음")
                
                # 수식 입력 (행 번호는 2부터 시작) - FCST 열 추가로 인해 한 칸씩 밀림
                row_num = last_row + 1
                if col_idx['판매 합계'] is not None:
                    data[col_idx['판매 합계']] = f'=K{row_num}*L{row_num}'
                if col_idx['원가계'] is not None:
                    data[col_idx['원가계']] = f'=N{row_num}*K{row_num}'
                if col_idx['이익액'] is not None:
                    data[col_idx['이익액']] = f'=M{row_num}-O{row_num}'
                if col_idx['이익율'] is not None:
                    data[col_idx['이익율']] = f'=IF(M{row_num}=0,0,P{row_num}/M{row_num})'
                
                ws.append_row(data, value_input_option='USER_ENTERED')
                
                # 숫자 컬럼들 오른쪽 정렬 설정 - FCST 열 추가로 인해 한 칸씩 밀림
                if col_idx['수량'] is not None:
                    ws.format(f'K{last_row}', {'horizontalAlignment': 'RIGHT'})
                if col_idx['판매 단가'] is not None:
                    ws.format(f'L{last_row}', {'horizontalAlignment': 'RIGHT'})
                if col_idx['판매 합계'] is not None:
                    ws.format(f'M{last_row}', {'horizontalAlignment': 'RIGHT'})
                if col_idx['원가'] is not None:
                    ws.format(f'N{last_row}', {'horizontalAlignment': 'RIGHT'})
                if col_idx['원가계'] is not None:
                    ws.format(f'O{last_row}', {'horizontalAlignment': 'RIGHT'})
                if col_idx['이익액'] is not None:
                    ws.format(f'P{last_row}', {'horizontalAlignment': 'RIGHT'})
                if col_idx['이익율'] is not None:
                    ws.format(f'Q{last_row}', {'horizontalAlignment': 'RIGHT'})
                
                last_row += 1
        except Exception as e:
            import traceback
            traceback.print_exc()
            from tkinter import messagebox
            messagebox.showerror('구글시트 저장 오류', str(e))
        
        # 저장 완료 후 부모 창 새로고침
        # self.refresh_parent_window() # 제거됨
        
        # Status가 80% 또는 100%일 때 renewal_list에 동기화
        selected_status = self.status_var.get()
        if selected_status in ['80%', '100%']:
            try:
                self.sync_to_renewal_list(quote_id, selected_status)
                print(f"renewal_list 동기화 완료: {quote_id}, Status: {selected_status}")
            except Exception as e:
                print(f"renewal_list 동기화 오류: {e}")
                import traceback
                traceback.print_exc()
        
        # 완료 메시지 표시 및 창 닫기
        from tkinter import messagebox
        result = messagebox.showinfo(
            "견적 작성 완료", 
            "견적서 리스트에 저장, 견적서 작성이 완료되었습니다",
            parent=self.win
        )
        
        # 견적 작성 창 닫기
        self.win.destroy()

    def refresh_parent_window(self):
        """부모 창 새로고침"""
        try:
            # 콜백 함수가 있으면 실행
            if self.callback and callable(self.callback):
                self.callback()
            # 콜백 함수가 없으면 기존 방식으로 시도
            elif hasattr(self.parent, 'load_data'):
                self.parent.load_data()
            elif hasattr(self.parent.master, 'load_data'):
                self.parent.master.load_data()
            elif hasattr(self.parent.master.master, 'load_data'):
                self.parent.master.master.load_data()
        except Exception as e:
            print(f"부모 창 새로고침 오류: {e}")
            pass

    def sync_to_renewal_list(self, quote_id, status):
        """
        Status가 80% 또는 100%일 때 renewal_list에 데이터 동기화
        - 견적서ID로 중복 체크: 이미 있으면 무시
        - 없으면 새로 추가
        """
        import gspread
        from google.oauth2.service_account import Credentials
        from config import SHEET_IDS, KEY_FILE, SCOPES
        
        print(f"=== renewal_list 동기화 시작 ===")
        print(f"견적서ID: {quote_id}, Status: {status}")
        
        try:
            # Google Sheets 인증
            creds = Credentials.from_service_account_file(KEY_FILE, scopes=SCOPES)
            gc = gspread.authorize(creds)
            
            # renewal_list 시트 열기
            sh = gc.open_by_key(SHEET_IDS['renewal_list'])
            try:
                ws = sh.worksheet('Sheet1')
            except:
                ws = sh.sheet1
            
            # 기존 데이터 로드하여 견적서ID 중복 체크
            all_data = ws.get_all_values()
            headers = all_data[0] if all_data else []
            
            # 견적서ID 컬럼 인덱스 찾기
            quote_id_col_idx = None
            for i, h in enumerate(headers):
                if h == '견적서ID':
                    quote_id_col_idx = i
                    break
            
            if quote_id_col_idx is None:
                print("경고: renewal_list에서 '견적서ID' 컬럼을 찾을 수 없습니다.")
                return
            
            # 이미 존재하는지 확인
            existing_quote_ids = [row[quote_id_col_idx] for row in all_data[1:] if len(row) > quote_id_col_idx]
            if quote_id in existing_quote_ids:
                print(f"견적서ID {quote_id}는 이미 renewal_list에 존재합니다. 동기화 건너뜀.")
                return
            
            # 컬럼 매핑 (renewal_list 컬럼명)
            renewal_cols = {h: i for i, h in enumerate(headers)}
            print(f"renewal_list 컬럼: {list(renewal_cols.keys())}")
            
            # quote_list 데이터를 renewal_list 형식으로 변환하여 추가
            for i, row_widgets in enumerate(self.quote_rows):
                # 데이터 준비
                new_row = [''] * len(headers)
                
                # 견적서ID
                if '견적서ID' in renewal_cols:
                    new_row[renewal_cols['견적서ID']] = quote_id
                
                # 고객사명
                if '고객사명' in renewal_cols:
                    customer = self.past_rows[0].get('고객사명', '') if self.past_rows else ''
                    new_row[renewal_cols['고객사명']] = customer
                
                # 거래처명
                if '거래처명' in renewal_cols:
                    client = self.past_rows[0].get('거래처명', '') if self.past_rows else ''
                    new_row[renewal_cols['거래처명']] = client
                
                # 대분류
                if '대분류' in renewal_cols:
                    category = row_widgets['대분류'].get()
                    new_row[renewal_cols['대분류']] = category
                
                # 소분류
                if '소분류' in renewal_cols:
                    subcategory = row_widgets['소분류'].get()
                    new_row[renewal_cols['소분류']] = subcategory
                
                # 제품
                if '제품' in renewal_cols:
                    product = row_widgets['제품'].get()
                    new_row[renewal_cols['제품']] = clean_product_name(product)
                
                # 수량
                if '수량' in renewal_cols:
                    qty = row_widgets['수량'].get()
                    new_row[renewal_cols['수량']] = qty
                
                # 판매 단가 / 판매단가
                for col_name in ['판매단가', '판매 단가']:
                    if col_name in renewal_cols:
                        unit_price = row_widgets['판매 단가'].get().replace(',', '')
                        new_row[renewal_cols[col_name]] = unit_price
                        break
                
                # 만료일
                if '만료일' in renewal_cols:
                    expiry = row_widgets['만료일'].get()
                    new_row[renewal_cols['만료일']] = expiry
                
                # 매입단가 (원가를 기본값으로 사용)
                for col_name in ['매입단가', '매입 단가', '실 매입 단가']:
                    if col_name in renewal_cols:
                        cost = row_widgets['원가'].get().replace(',', '')
                        new_row[renewal_cols[col_name]] = cost
                        break
                
                # 매입처
                if '매입처' in renewal_cols:
                    supplier = row_widgets.get('매입처', {})
                    if hasattr(supplier, 'get'):
                        new_row[renewal_cols['매입처']] = supplier.get()
                
                # 영업사원
                if '영업사원' in renewal_cols:
                    new_row[renewal_cols['영업사원']] = self.username
                
                # 유동유무
                if '유동유무' in renewal_cols:
                    new_row[renewal_cols['유동유무']] = '유동'
                
                # 계산서 발행일 (100%일 때만)
                if status == '100%' and '계산서 발행일' in renewal_cols:
                    from datetime import datetime
                    new_row[renewal_cols['계산서 발행일']] = datetime.today().strftime('%Y-%m-%d')
                
                # 행 추가
                ws.append_row(new_row, value_input_option='USER_ENTERED')
                print(f"renewal_list에 행 추가 완료: {quote_id}, 제품: {new_row[renewal_cols.get('제품', 0)] if '제품' in renewal_cols else 'N/A'}")
            
            print(f"=== renewal_list 동기화 완료 ===")
            
        except Exception as e:
            print(f"renewal_list 동기화 오류: {e}")
            import traceback
            traceback.print_exc()
            raise

    def create_exchange_tables_side_by_side(self):
        """환율 정보와 가격표 환율을 나란히 배치"""
        # 환율 정보와 가격표 환율을 담을 컨테이너 프레임
        container_frame = tk.Frame(self.win)
        container_frame.pack(pady=(10, 10), padx=10, fill='x')
        
        # 환율 정보 (왼쪽)
        exchange_frame = tk.LabelFrame(container_frame, text="환율 정보", font=("맑은 고딕", 9, "bold"), bg="white", relief="solid", bd=1, labelanchor='nw')
        exchange_frame.pack(side='left', fill='both', expand=True, padx=(0, 5))
        
        headers = ['구분', '환율']
        for i, h in enumerate(headers):
            tk.Label(exchange_frame, text=h, relief='solid', width=16, bg='#F1F5F9', fg="#374151", font=("맑은 고딕", 10, "bold"), anchor='center').grid(row=0, column=i, sticky='nsew')
        
        usd_rate = self.get_naver_fx_rate('USD')
        eur_rate = self.get_naver_fx_rate('EUR')
        tk.Label(exchange_frame, text='USD', relief='solid', width=16, bg='white', anchor='center').grid(row=1, column=0, sticky='nsew')
        tk.Label(exchange_frame, text=usd_rate, relief='solid', width=16, bg='white', anchor='e').grid(row=1, column=1, sticky='nsew')
        tk.Label(exchange_frame, text='EURO', relief='solid', width=16, bg='white', anchor='center').grid(row=2, column=0, sticky='nsew')
        tk.Label(exchange_frame, text=eur_rate, relief='solid', width=16, bg='white', anchor='e').grid(row=2, column=1, sticky='nsew')
        
        # 가격표 환율 (오른쪽)
        price_fx_frame = tk.LabelFrame(container_frame, text="가격표 환율", font=("맑은 고딕", 9, "bold"), bg="white", relief="solid", bd=1, labelanchor='nw')
        price_fx_frame.pack(side='right', fill='both', expand=True, padx=(5, 0))
        
        headers = ['구분', '환율', '적용 환율', '']
        for i, h in enumerate(headers):
            tk.Label(price_fx_frame, text=h, relief='solid', width=12, bg='#F1F5F9', fg="#374151", font=("맑은 고딕", 10, "bold"), anchor='center').grid(row=0, column=i, sticky='nsew')
        
        usd_rate, eur_rate = self.get_price_fx_values()
        self.usd_apply_var = tk.StringVar()
        self.eur_apply_var = tk.StringVar()
        
        # USD 행
        tk.Label(price_fx_frame, text='USD', relief='solid', width=12, bg='white', anchor='center').grid(row=1, column=0, sticky='nsew')
        self.usd_rate_var = tk.StringVar(value=usd_rate)
        tk.Entry(price_fx_frame, textvariable=self.usd_rate_var, width=12, state='readonly', justify='right', relief='solid', font=("맑은 고딕", 10), bg='white').grid(row=1, column=1, sticky='nsew')
        usd_apply_entry = tk.Entry(price_fx_frame, textvariable=self.usd_apply_var, width=12, justify='right', relief='solid', font=("맑은 고딕", 10), bg='white')
        usd_apply_entry.grid(row=1, column=2, sticky='nsew')
        tk.Button(price_fx_frame, text='적용', width=6, command=lambda: self.apply_price_fx('USD'), font=("맑은 고딕", 10, "bold"), bg="#10B981", fg="white", relief="flat").grid(row=1, column=3, sticky='nsew')
        
        # EUR 행
        tk.Label(price_fx_frame, text='EURO', relief='solid', width=12, bg='white', anchor='center').grid(row=2, column=0, sticky='nsew')
        self.eur_rate_var = tk.StringVar(value=eur_rate)
        tk.Entry(price_fx_frame, textvariable=self.eur_rate_var, width=12, state='readonly', justify='right', relief='solid', font=("맑은 고딕", 10), bg='white').grid(row=2, column=1, sticky='nsew')
        eur_apply_entry = tk.Entry(price_fx_frame, textvariable=self.eur_apply_var, width=12, justify='right', relief='solid', font=("맑은 고딕", 10), bg='white')
        eur_apply_entry.grid(row=2, column=2, sticky='nsew')
        tk.Button(price_fx_frame, text='적용', width=6, command=lambda: self.apply_price_fx('EUR'), font=("맑은 고딕", 10, "bold"), bg="#10B981", fg="white", relief="flat").grid(row=2, column=3, sticky='nsew')

    def get_naver_fx_rate(self, currency):
        """네이버 환율 리스트에서 USD, EUR 매매기준율만 크롤링"""
        try:
            import requests
            from bs4 import BeautifulSoup
            url = 'https://finance.naver.com/marketindex/exchangeList.naver'
            resp = requests.get(url, timeout=3)
            soup = BeautifulSoup(resp.text, 'html.parser')
            rows = soup.select('table.tbl_exchange tbody tr')
            for row in rows:
                tds = row.find_all('td')
                if len(tds) >= 2:
                    name = tds[0].get_text(strip=True)
                    rate = tds[1].get_text(strip=True).replace(',', '')
                    if currency == 'USD' and 'USD' in name:
                        return f'{float(rate):,.2f}'
                    if currency == 'EUR' and ('EUR' in name or '유로' in name):
                        return f'{float(rate):,.2f}'
            return ''
        except Exception as e:
            print(f'환율 정보 오류: {e}')
            # 기본값 반환
            if currency == 'USD':
                return '1,300.00'
            elif currency == 'EUR':
                return '1,400.00'
            return ''

    def get_product_info_from_price_list(self, product_name):
        """price_list에서 제품 정보 가져오기 (캐시 최적화)"""
        try:
            # 캐시 키 생성
            cache_key = f"product_info_{hashlib.md5(product_name.encode()).hexdigest()}"
            
            # 캐시에서 먼저 확인
            from core.data_loader import get_cached_data, set_cached_data
            cached_info = get_cached_data(cache_key)
            if cached_info is not None:
                print(f"제품 정보를 캐시에서 로드: {product_name}")
                return cached_info
            
            # API 호출 최소화를 위해 safe_api_call 사용
            gc = get_global_auth()
            sh = gc.open_by_key('1OGt__Olrempc5ZUy6QqqhvgRseFLanv1xCPaw5Y2qJk')
            ws = sh.worksheet('Sheet1')
            
            # 모든 데이터 가져오기 (안전한 API 호출)
            from core.data_loader import safe_api_call
            all_data = safe_api_call(lambda: ws.get_all_records())
            
            if all_data:
                print(f"제품 검색 중: '{product_name}'")
                print(f"price_list 데이터 개수: {len(all_data)}")
                
                # 1. 제품명 정확 일치 검색
                for i, row in enumerate(all_data):
                    if '제품' in row and row['제품'] == product_name:
                        product_info = {
                            '구분': row.get('구분', ''),
                            'USD': row.get('USD', ''),
                            'EUR': row.get('EUR', ''),
                            'DC': row.get('DC', ''),
                            '매입처': row.get('매입처', '')
                        }
                        print(f"제품 정보 찾음 (정확 일치, 행 {i+1}): {product_info}")
                        # 캐시에 저장
                        set_cached_data(cache_key, product_info)
                        return product_info
                
                # 2. 정확 일치 실패 시 부분 일치 검색
                product_name_lower = product_name.lower().strip()
                for i, row in enumerate(all_data):
                    if '제품' in row:
                        row_product = str(row['제품']).lower().strip()
                        if product_name_lower in row_product or row_product in product_name_lower:
                            product_info = {
                                '구분': row.get('구분', ''),
                                'USD': row.get('USD', ''),
                                'EUR': row.get('EUR', ''),
                                'DC': row.get('DC', ''),
                                '매입처': row.get('매입처', '')
                            }
                            print(f"제품 정보 찾음 (부분 일치, 행 {i+1}): {product_info}")
                            # 캐시에 저장
                            set_cached_data(cache_key, product_info)
                            return product_info
                
                print(f"제품 '{product_name}'을 찾을 수 없습니다.")
                # 디버깅을 위해 처음 몇 개 제품명 출력
                for i, row in enumerate(all_data[:5]):
                    if '제품' in row:
                        print(f"  행 {i+1}: '{row['제품']}'")
            
            # 제품을 찾지 못한 경우 빈 정보를 캐시에 저장 (API 호출 방지)
            empty_info = {'구분': '', 'USD': '', 'EUR': '', 'DC': '', '매입처': ''}
            set_cached_data(cache_key, empty_info)
            return None
            
        except Exception as e:
            print(f"price_list에서 제품 정보 읽기 오류: {e}")
            # 오류 발생 시 빈 정보 반환 (API 호출 방지)
            return {'구분': '', 'USD': '', 'EUR': '', 'DC': '', '매입처': ''}

    def get_price_fx_values(self):
        # price_list Sheet2 A2(USD), B2(EUR) 값 읽기
        try:
            gc = get_global_auth()
            
            def get_usd():
                sh = gc.open_by_key('1OGt__Olrempc5ZUy6QqqhvgRseFLanv1xCPaw5Y2qJk')
                ws = sh.worksheet('Sheet2')
                return ws.acell('A2').value
            
            def get_eur():
                sh = gc.open_by_key('1OGt__Olrempc5ZUy6QqqhvgRseFLanv1xCPaw5Y2qJk')
                ws = sh.worksheet('Sheet2')
                return ws.acell('B2').value
            
            usd = safe_api_call(get_usd)
            eur = safe_api_call(get_eur)
            
            # 값이 None이거나 빈 문자열인 경우 기본값 사용
            if not usd or usd == '' or usd == 'None':
                usd = '1,390.00'  # 기본 USD 환율
                print("견적 작성(기존): USD 환율이 비어있어 기본값을 사용합니다.")
            else:
                usd = self.format_fx_value(usd)
                
            if not eur or eur == '' or eur == 'None':
                eur = '1,630.00'  # 기본 EUR 환율
                print("견적 작성(기존): EUR 환율이 비어있어 기본값을 사용합니다.")
            else:
                eur = self.format_fx_value(eur)
            
            print(f"견적 작성(기존): 환율 데이터 로드 완료 - USD: {usd}, EUR: {eur}")
            return usd, eur
        except Exception as e:
            print(f'견적 작성(기존): 가격표 환율 읽기 오류 - {e}')
            print("견적 작성(기존): 기본값으로 계속 진행합니다.")
            return '1,390.00', '1,630.00'  # 기본값 반환

    def format_fx_value(self, val):
        try:
            if val is None or val == '' or str(val).strip() == '':
                return ''
            # 쉼표 제거 후 숫자로 변환
            clean_val = str(val).replace(',', '').strip()
            fval = float(clean_val)
            return f"{fval:,.2f}"
        except (ValueError, TypeError) as e:
            print(f"견적 작성(기존): 환율 값 형식 변환 오류 - {val}, 오류: {e}")
            return str(val) if val else ''



    def apply_price_fx(self, currency=None):
        # 적용 환율을 price_list Sheet2 A2(USD), B2(EUR)에 업데이트
        try:
            gc = get_global_auth()
            
            def update_usd():
                sh = gc.open_by_key('1OGt__Olrempc5ZUy6QqqhvgRseFLanv1xCPaw5Y2qJk')
                ws = sh.worksheet('Sheet2')
                value = self.usd_apply_var.get().strip()
                if not value:
                    print("견적 작성(기존): USD 환율 값이 비어있습니다.")
                    return
                try:
                    # 숫자로 변환하여 유효성 검사
                    num_value = float(value.replace(',', ''))
                    value_fmt = f"{num_value:,.2f}"
                    ws.update_acell('A2', value_fmt)
                    self.usd_rate_var.set(value_fmt)
                    self.usd_apply_var.set(value_fmt)
                    print(f"견적 작성(기존): USD 환율 업데이트 완료 - {value_fmt}")
                except ValueError as ve:
                    print(f"견적 작성(기존): USD 환율 값 형식 오류 - {value}")
                    messagebox.showerror("오류", f"USD 환율 값이 올바르지 않습니다: {value}")
            
            def update_eur():
                sh = gc.open_by_key('1OGt__Olrempc5ZUy6QqqhvgRseFLanv1xCPaw5Y2qJk')
                ws = sh.worksheet('Sheet2')
                value = self.eur_apply_var.get().strip()
                if not value:
                    print("견적 작성(기존): EUR 환율 값이 비어있습니다.")
                    return
                try:
                    # 숫자로 변환하여 유효성 검사
                    num_value = float(value.replace(',', ''))
                    value_fmt = f"{num_value:,.2f}"
                    ws.update_acell('B2', value_fmt)
                    self.eur_rate_var.set(value_fmt)
                    self.eur_apply_var.set(value_fmt)
                    print(f"견적 작성(기존): EUR 환율 업데이트 완료 - {value_fmt}")
                except ValueError as ve:
                    print(f"견적 작성(기존): EUR 환율 값 형식 오류 - {value}")
                    messagebox.showerror("오류", f"EUR 환율 값이 올바르지 않습니다: {value}")
            
            if currency == 'USD':
                safe_api_call(update_usd)
                print("견적 작성(기존): USD 환율 적용 완료")
            elif currency == 'EUR':
                safe_api_call(update_eur)
                print("견적 작성(기존): EUR 환율 적용 완료")
            else:
                print(f"견적 작성(기존): 알 수 없는 통화 - {currency}")
                return
                
            # 환율 적용 후 price_list 데이터 새로고침 및 견적 행 원가 갱신
            self.load_price_data()
            for row in self.quote_rows:
                # 환율 변경 시 기존 원가/매입처 값을 유지하면서 재계산
                self.on_product_change(row, preserve_existing=True)
        except Exception as e:
            print(f'견적 작성(기존): 가격표 환율 적용 오류 - {e}')
            messagebox.showerror("오류", f"환율 적용 중 오류가 발생했습니다:\n{e}")

    def select_row(self, row_index):
        """행 선택"""
        try:
            # 이전 선택 해제
            if hasattr(self, 'selected_row_index') and self.selected_row_index is not None and self.selected_row_index < len(self.quote_rows):
                self.clear_row_selection(self.selected_row_index)
            
            # 새로운 행 선택
            self.selected_row_index = row_index
            if row_index < len(self.quote_rows):
                self.highlight_row_selection(row_index)
                print(f"견적 작성(기존): 행 {row_index + 1} 선택됨")
            else:
                print(f"견적 작성(기존): 잘못된 행 인덱스 {row_index} (총 {len(self.quote_rows)}행)")
                
        except Exception as e:
            print(f"견적 작성(기존): 행 선택 오류 - {e}")
            self.selected_row_index = None

    def clear_row_selection(self, row_index):
        """행 선택 해제 (시각적 피드백 제거)"""
        try:
            if row_index >= len(self.quote_rows):
                return
                
            row_widgets = self.quote_rows[row_index]
            row_num = row_index + 1
            
            col_keys = ['대분류', '제품', '수량', '판매 단가', '판매 합계', '만료일', '원가', '원가계', '이익액', '이익율']
            for col_idx, col_name in enumerate(col_keys):
                widget = self.get_widget_by_col(row_widgets, col_idx)
                if widget:
                    try:
                        # ttk.Combobox는 기본 스타일로 복원
                        if hasattr(widget, 'configure') and 'style' in widget.configure():
                            widget.configure(style='TCombobox')
                        else:
                            widget.configure(bg='white')
                    except:
                        # 배경색 변경이 실패하면 무시
                        pass
        except Exception as e:
            print(f"견적 작성(기존): 행 선택 해제 오류 - {e}")

    def highlight_row_selection(self, row_index):
        """행 선택 시 시각적 피드백 (배경색 변경)"""
        try:
            if row_index >= len(self.quote_rows):
                return
                
            row_widgets = self.quote_rows[row_index]
            
            col_keys = ['대분류', '제품', '수량', '판매 단가', '판매 합계', '만료일', '원가', '원가계', '이익액', '이익율']
            for col_idx, col_name in enumerate(col_keys):
                widget = self.get_widget_by_col(row_widgets, col_idx)
                if widget:
                    try:
                        # ttk.Combobox는 style을 통해 배경색 변경
                        if hasattr(widget, 'configure') and 'style' in widget.configure():
                            # 커스텀 스타일 생성 (파란색 배경)
                            style = ttk.Style()
                            style.configure('Selected.TCombobox', fieldbackground='#E0F2FE')
                            widget.configure(style='Selected.TCombobox')
                        else:
                            widget.configure(bg='#E0F2FE')  # 연한 파란색 배경
                    except:
                        # 배경색 변경이 실패하면 무시
                        pass
        except Exception as e:
            print(f"견적 작성(기존): 행 하이라이트 오류 - {e}")

    def delete_selected_row(self):
        """선택된 행 삭제"""
        try:
            if not hasattr(self, 'selected_row_index') or self.selected_row_index is None:
                from tkinter import messagebox
                messagebox.showwarning("선택 필요", "삭제할 행을 먼저 선택하세요.")
                return
            
            if len(self.quote_rows) <= 1:
                from tkinter import messagebox
                messagebox.showwarning("삭제 불가", "최소 1행은 유지해야 합니다.")
                return
            
            # 선택된 행 번호
            row_num = self.selected_row_index + 1
            
            # 선택된 행의 위젯들 제거
            widgets_to_destroy = []
            for widget in self.table_frame.winfo_children():
                try:
                    grid_info = widget.grid_info()
                    if grid_info and grid_info['row'] == row_num:
                        widgets_to_destroy.append(widget)
                except:
                    continue
            
            for widget in widgets_to_destroy:
                try:
                    widget.destroy()
                except:
                    continue
            
            # quote_rows에서 제거
            self.quote_rows.pop(self.selected_row_index)
            
            # 나머지 행들의 인덱스 업데이트
            for i, row in enumerate(self.quote_rows):
                row['row_index'] = i
            
            # 위젯 재배치
            self.rearrange_widgets()
            
            # 선택 상태 초기화
            self.selected_row_index = None
            print(f"견적 작성(기존): 행 {row_num} 삭제 완료")
            
        except Exception as e:
            print(f"견적 작성(기존): 행 삭제 오류 - {e}")
            from tkinter import messagebox
            messagebox.showerror("오류", f"행 삭제 중 오류가 발생했습니다:\n{e}")

    def rearrange_widgets(self):
        """위젯들을 새로운 행 순서에 맞게 재배치"""
        try:
            # 실제 컬럼 순서 (get_widget_by_col과 일치)
            col_keys = ['대분류', '소분류', '제품', '수량', '판매 단가', '판매 합계', '만료일', '원가', '원가계', '이익액', '이익율', '매입처']
            
            for row_idx, row_widgets in enumerate(self.quote_rows):
                row_num = row_idx + 1
                
                # 모든 컬럼들 재배치
                for col_idx, col_key in enumerate(col_keys):
                    try:
                        widget = self.get_widget_by_col(row_widgets, col_idx)
                        if widget and hasattr(widget, 'grid'):  # 위젯이 존재하고 grid 메서드가 있는 경우만
                            widget.grid(row=row_num, column=col_idx, sticky='ew', padx=1, pady=1)
                    except Exception as e:
                        print(f"견적 작성(기존): 컬럼 {col_key} 재배치 오류 - {e}")
                        continue
                
                # 행 선택 이벤트 바인딩 재설정
                for key, widget in row_widgets.items():
                    try:
                        if hasattr(widget, 'bind') and key.endswith(('_combo', '_entry')):
                            # 기존 바인딩 제거
                            widget.unbind('<Button-1>')
                            # 새로운 바인딩 추가
                            widget.bind('<Button-1>', lambda e, idx=row_idx: self.select_row(idx))
                    except Exception as e:
                        print(f"견적 작성(기존): 이벤트 바인딩 오류 - {e}")
                        continue
            
            print("견적 작성(기존): 위젯 재배치 완료")
            
        except Exception as e:
            print(f"견적 작성(기존): 위젯 재배치 오류 - {e}")
            # 오류 발생 시 전체 테이블을 다시 생성
            self.refresh_table() 

    def _get_next_quote_number(self, date_str):
        """오늘 날짜의 전체 견적서 번호 생성"""
        try:
            import gspread
            from google.oauth2.service_account import Credentials
            creds = Credentials.from_service_account_file(KEY_FILE, scopes=SCOPES)
            gc = gspread.authorize(creds)
            sh = gc.open_by_key('1xjIjFe1Q9dq2zeQOkBSpbqIUOyXAxbyTO8lKN-r1lw0')
            ws = sh.worksheet('Sheet1')
            
            # 전체 데이터 가져오기
            all_data = ws.get_all_values()
            if len(all_data) <= 1:  # 헤더만 있거나 비어있음
                return 1
            
            # 컬럼 인덱스 찾기 (ID 또는 견적서ID 컬럼 찾기)
            header = all_data[0]
            id_idx = None
            date_idx = None
            
            for idx, h in enumerate(header):
                if h == 'ID':
                    id_idx = idx
                elif h == '견적서ID':
                    id_idx = idx  # 견적서ID도 ID로 사용
                elif h == '일자':
                    date_idx = idx
            
            if id_idx is None or date_idx is None:
                return 1
            
            # 오늘 날짜의 모든 견적서 번호 찾기
            max_number = 0
            found_ids = set()  # 중복 제거를 위한 set
            
            for row in all_data[1:]:  # 헤더 제외
                if len(row) > max(id_idx, date_idx):
                    row_date = row[date_idx]
                    row_id = row[id_idx]
                    
                    if (row_date.startswith(date_str) and 
                        row_id and 
                        '_' in row_id and
                        row_id not in found_ids):  # 중복 제거
                        try:
                            # ID에서 번호 추출 (예: 2025-08-20_001 -> 1)
                            number_part = row_id.split('_')[-1]
                            if number_part.isdigit():
                                max_number = max(max_number, int(number_part))
                                found_ids.add(row_id)  # 중복 제거를 위해 추가
                                print(f"견적서ID 발견: {row_id} -> 번호: {number_part}")
                        except:
                            continue
            
            next_number = max_number + 1
            print(f"오늘 날짜({date_str}) 최대 번호: {max_number}, 다음 번호: {next_number}")
            return next_number
            
        except Exception as e:
            print(f"견적서 번호 생성 오류: {e}")
            return 1

    def recreate_table(self):
        """테이블 전체를 재생성 (행 삭제 후 안전한 재배치)"""
        try:
            # 기존 테이블 위젯들 모두 제거 (헤더 제외)
            widgets_to_destroy = []
            for widget in self.table_frame.winfo_children():
                grid_info = widget.grid_info()
                if grid_info and grid_info['row'] > 0:  # 헤더(0행) 제외
                    widgets_to_destroy.append(widget)
            
            for widget in widgets_to_destroy:
                widget.destroy()
            
            # 기존 데이터로 테이블 재생성
            for row_idx, row_data in enumerate(self.quote_rows):
                row_num = row_idx + 1
                
                # 각 컬럼별로 위젯 재생성
                col_keys = ['대분류', '소분류', '제품', '수량', '판매 단가', '판매 합계', '만료일', '원가', '원가계', '이익액', '이익율', '매입처']
                
                for col_idx, col_key in enumerate(col_keys):
                    # 기존 위젯이 있으면 재사용, 없으면 새로 생성
                    if col_key in row_data:
                        if col_key in ['대분류', '소분류', '제품']:
                            # Combobox 재생성
                            combo_key = f'{col_key}_combo'
                            if combo_key in row_data:
                                combo = row_data[combo_key]
                                combo.grid(row=row_num, column=col_idx, sticky='ew', padx=1, pady=1)
                        else:
                            # Entry 재생성
                            entry_key = f'{col_key}_entry'
                            if entry_key in row_data:
                                entry = row_data[entry_key]
                                entry.grid(row=row_num, column=col_idx, sticky='ew', padx=1, pady=1)
                
                # 행 선택 이벤트 바인딩 재설정
                for key, widget in row_data.items():
                    if hasattr(widget, 'bind') and key.endswith(('_combo', '_entry')):
                        # 기존 바인딩 제거
                        widget.unbind('<Button-1>')
                        # 새로운 바인딩 추가
                        widget.bind('<Button-1>', lambda e, idx=row_idx: self.select_row(idx))
            
            print("견적 작성(기존): 테이블 재생성 완료")
            
        except Exception as e:
            print(f"견적 작성(기존): 테이블 재생성 오류 - {e}")
            # 오류 발생 시 전체 테이블을 다시 생성
            self.refresh_table()

    def refresh_table(self):
        """테이블 전체를 새로 생성 (긴급 복구용)"""
        try:
            # 기존 데이터 백업
            backup_data = []
            for row in self.quote_rows:
                row_data = {}
                for key, widget in row.items():
                    if key == 'row_index':
                        row_data[key] = widget
                    elif hasattr(widget, 'get'):
                        row_data[key] = widget.get()
                backup_data.append(row_data)
            
            # 테이블 프레임 초기화
            for widget in self.table_frame.winfo_children():
                widget.destroy()
            
            # 헤더 재생성
            self.create_quote_headers()
            
            # 데이터로 테이블 재생성
            self.quote_rows = []
            for row_data in backup_data:
                self.add_row_with_data(row_data)
            
            print("견적 작성(기존): 테이블 전체 새로고침 완료")
            
        except Exception as e:
            print(f"견적 작성(기존): 테이블 새로고침 오류 - {e}")

    def add_row_with_data(self, row_data):
        """기존 데이터로 행 추가"""
        try:
            row_num = len(self.quote_rows) + 1
            row_widgets = {}
            row_widgets['row_index'] = len(self.quote_rows)
            
            # 대분류
            category_var = tk.StringVar(value=row_data.get('대분류', ''))
            category_combo = ttk.Combobox(self.table_frame, textvariable=category_var, 
                                        values=self.categories, width=self.col_widths[0], 
                                        state='normal', font=("맑은 고딕", 10))
            category_combo.grid(row=row_num, column=0, sticky='ew', padx=1, pady=1)
            row_widgets['대분류'] = category_var
            row_widgets['대분류_combo'] = category_combo
            category_combo.bind('<<ComboboxSelected>>', lambda e, row=row_widgets: self.on_category_change(row))
            category_combo.bind('<Return>', lambda e, row=row_widgets: self.on_category_change(row))
            
            # 소분류
            subcategory_var = tk.StringVar(value=row_data.get('소분류', ''))
            subcategory_combo = ttk.Combobox(self.table_frame, textvariable=subcategory_var, 
                                           width=self.col_widths[1], state='normal', 
                                           font=("맑은 고딕", 10))
            subcategory_combo.grid(row=row_num, column=1, sticky='ew', padx=1, pady=1)
            row_widgets['소분류'] = subcategory_var
            row_widgets['소분류_combo'] = subcategory_combo
            subcategory_combo.bind('<<ComboboxSelected>>', lambda e, row=row_widgets: self.on_subcategory_change(row))
            subcategory_combo.bind('<Return>', lambda e, row=row_widgets: self.on_subcategory_change(row))
            
            # 제품
            product_var = tk.StringVar(value=row_data.get('제품', ''))
            product_combo = ttk.Combobox(self.table_frame, textvariable=product_var, 
                                       width=self.col_widths[2], state='normal', 
                                       font=("맑은 고딕", 10))
            product_combo.grid(row=row_num, column=2, sticky='ew', padx=1, pady=1)
            row_widgets['제품'] = product_var
            row_widgets['제품_combo'] = product_combo
            product_combo.bind('<<ComboboxSelected>>', lambda e, row=row_widgets: self.on_product_change(row))
            product_combo.bind('<Return>', lambda e, row=row_widgets: self.on_product_search_and_apply(row))
            
            # 수량
            qty_var = tk.StringVar(value=row_data.get('수량', ''))
            qty_entry = tk.Entry(self.table_frame, textvariable=qty_var, 
                               width=self.col_widths[3], justify='center', 
                               font=("맑은 고딕", 10))
            qty_entry.grid(row=row_num, column=3, sticky='ew', padx=1, pady=1)
            qty_entry.bind('<KeyRelease>', lambda e, row=row_widgets: self.on_qty_change(row))
            qty_entry.bind('<FocusOut>', lambda e, row=row_widgets: self.on_qty_change(row))
            row_widgets['수량'] = qty_var
            row_widgets['수량_entry'] = qty_entry
            
            # 판매 단가
            price_var = tk.StringVar(value=row_data.get('판매 단가', ''))
            price_entry = tk.Entry(self.table_frame, textvariable=price_var, 
                                 width=self.col_widths[4], justify='right', 
                                 font=("맑은 고딕", 10))
            price_entry.grid(row=row_num, column=4, sticky='ew', padx=1, pady=1)
            price_entry.bind('<KeyRelease>', lambda e, row=row_widgets: self.on_price_change(row))
            price_entry.bind('<FocusOut>', lambda e, row=row_widgets: self.on_price_change(row))
            row_widgets['판매 단가'] = price_var
            row_widgets['판매 단가_entry'] = price_entry
            
            # 판매 합계
            total_var = tk.StringVar(value=row_data.get('판매 합계', ''))
            total_entry = tk.Entry(self.table_frame, textvariable=total_var, 
                                 state='readonly', width=self.col_widths[5], 
                                 justify='right', font=("맑은 고딕", 10))
            total_entry.grid(row=row_num, column=5, sticky='ew', padx=1, pady=1)
            row_widgets['판매 합계'] = total_var
            row_widgets['판매 합계_entry'] = total_entry
            
            # 만료일
            period_var = tk.StringVar(value=row_data.get('만료일', ''))
            period_entry = tk.Entry(self.table_frame, textvariable=period_var, 
                                  width=self.col_widths[6], font=("맑은 고딕", 10))
            period_entry.grid(row=row_num, column=6, sticky='ew', padx=1, pady=1)
            row_widgets['만료일'] = period_var
            row_widgets['만료일_entry'] = period_entry
            
            # 원가
            cost_var = tk.StringVar(value=row_data.get('원가', ''))
            cost_entry = tk.Entry(self.table_frame, textvariable=cost_var, 
                                width=self.col_widths[7], justify='right', 
                                font=("맑은 고딕", 10))
            cost_entry.grid(row=row_num, column=7, sticky='ew', padx=1, pady=1)
            row_widgets['원가'] = cost_var
            row_widgets['원가_entry'] = cost_entry
            
            # 원가계
            cost_total_var = tk.StringVar(value=row_data.get('원가계', ''))
            cost_total_entry = tk.Entry(self.table_frame, textvariable=cost_total_var, 
                                      state='readonly', width=self.col_widths[8], 
                                      justify='right', font=("맑은 고딕", 10))
            cost_total_entry.grid(row=row_num, column=8, sticky='ew', padx=1, pady=1)
            row_widgets['원가계'] = cost_total_var
            row_widgets['원가계_entry'] = cost_total_entry
            
            # 이익액
            profit_var = tk.StringVar(value=row_data.get('이익액', ''))
            profit_entry = tk.Entry(self.table_frame, textvariable=profit_var, 
                                  state='readonly', width=self.col_widths[9], 
                                  justify='right', font=("맑은 고딕", 10))
            profit_entry.grid(row=row_num, column=9, sticky='ew', padx=1, pady=1)
            row_widgets['이익액'] = profit_var
            row_widgets['이익액_entry'] = profit_entry
            
            # 이익율
            profit_rate_var = tk.StringVar(value=row_data.get('이익율', ''))
            profit_rate_entry = tk.Entry(self.table_frame, textvariable=profit_rate_var, 
                                       state='readonly', width=self.col_widths[10], 
                                       justify='right', font=("맑은 고딕", 10))
            profit_rate_entry.grid(row=row_num, column=10, sticky='ew', padx=1, pady=1)
            row_widgets['이익율'] = profit_rate_var
            row_widgets['이익율_entry'] = profit_rate_entry
            
            # 매입처
            supplier_var = tk.StringVar(value=row_data.get('매입처', ''))
            supplier_entry = tk.Entry(self.table_frame, textvariable=supplier_var, 
                                    width=self.col_widths[11], font=("맑은 고딕", 10))
            supplier_entry.grid(row=row_num, column=11, sticky='ew', padx=1, pady=1)
            row_widgets['매입처'] = supplier_var
            row_widgets['매입처_entry'] = supplier_entry
            
            # 행 선택 이벤트 바인딩
            for key, widget in row_widgets.items():
                if hasattr(widget, 'bind') and key.endswith(('_combo', '_entry')):
                    widget.bind('<Button-1>', lambda e, idx=len(self.quote_rows): self.select_row(idx))
            
            self.quote_rows.append(row_widgets)
            
        except Exception as e:
            print(f"견적 작성(기존): 행 데이터 추가 오류 - {e}")