import os
import sys
import shutil
import datetime
import tkinter as tk
from tkinter import messagebox, scrolledtext
import openpyxl
import win32com.client as wc  # pywin32 필요
from docx import Document  # python-docx 필요
import re

def resource_path(relative_path):
    import sys, os
    if getattr(sys, 'frozen', False):
        base_path = os.path.dirname(sys.executable)
    else:
        base_path = os.path.abspath('.')
    return os.path.join(base_path, relative_path)

def resource_root():
    import sys, os
    if getattr(sys, 'frozen', False):
        return os.path.dirname(sys.executable)
    else:
        return os.path.abspath('.')

def parse_pasted_text(pasted_text):
    lines = pasted_text.splitlines()
    result = {}
    for line in lines:
        line = line.strip().lstrip('•').strip()
        if ':' in line:
            key, value = line.split(':', 1)
            result[key.strip().lower()] = value.strip()
    return result

def create_certificate(company, product, pasted_text, license_type):
    base_dir = os.path.dirname(os.path.abspath(__file__))
    template_path = resource_path('license_certificate/BorisFX/BorisFX_origin.docx')
    output_dir = os.path.join(resource_root(), 'license_certificate', 'BorisFX')
    os.makedirs(output_dir, exist_ok=True)
    doc = Document(template_path)
    today = datetime.date.today()
    formatted_today = today.strftime("%B %d, %Y")
    parsed = parse_pasted_text(pasted_text)
    host = parsed.get('host', '')
    quantity = parsed.get('quantity', '')
    exp_str = parsed.get('expiration', '')
    serial = parsed.get('activation key #', '') or parsed.get('activation key', '')
    has_subscription = 'annual subscription' in pasted_text.lower()
    annual_subscription = 'Annual Subscription' if has_subscription else ''
    exp_formatted = ''
    if exp_str:
        try:
            exp_date = datetime.datetime.strptime(exp_str, '%b %d, %Y').date()
        except ValueError:
            try:
                exp_date = datetime.datetime.strptime(exp_str, '%B %d, %Y').date()
            except Exception:
                exp_date = None
        if exp_date:
            exp_formatted = exp_date.strftime('%Y/%m/%d')
    def replace_text_in_paragraph(paragraph, old, new):
        full_text = ''.join(run.text for run in paragraph.runs)
        if old in full_text:
            full_text = full_text.replace(old, new)
            for run in paragraph.runs:
                run.text = ''
            if paragraph.runs:
                paragraph.runs[0].text = full_text
    for para in doc.paragraphs:
        replace_text_in_paragraph(para, '{Today}', formatted_today)
        replace_text_in_paragraph(para, '{Company Name}', company)
        replace_text_in_paragraph(para, '{Product}', product)
        replace_text_in_paragraph(para, '{Host}', host)
        replace_text_in_paragraph(para, '{Qty}', quantity)
        replace_text_in_paragraph(para, '{Serial}', serial)
        replace_text_in_paragraph(para, '{Date}', exp_formatted)
        replace_text_in_paragraph(para, '{Annual Subscription}', annual_subscription)
        replace_text_in_paragraph(para, '{License type}', license_type)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    replace_text_in_paragraph(para, '{Today}', formatted_today)
                    replace_text_in_paragraph(para, '{Company Name}', company)
                    replace_text_in_paragraph(para, '{Product}', product)
                    replace_text_in_paragraph(para, '{Host}', host)
                    replace_text_in_paragraph(para, '{Qty}', quantity)
                    replace_text_in_paragraph(para, '{Serial}', serial)
                    replace_text_in_paragraph(para, '{Date}', exp_formatted)
                    replace_text_in_paragraph(para, '{Annual Subscription}', annual_subscription)
                    replace_text_in_paragraph(para, '{License type}', license_type)
    date_str = today.strftime("%Y%m%d")
    safe_company = re.sub(r'\s+', '', company)
    safe_product = re.sub(r'\s+', '', product)
    filename_base = f"BorisFX_{safe_product}_{safe_company}_{date_str}"
    new_docx_path = os.path.join(output_dir, filename_base + '.docx')
    doc.save(new_docx_path)
    word = wc.Dispatch('Word.Application')
    word.Visible = False
    docx_abs = os.path.abspath(new_docx_path)
    pdf_path = os.path.join(output_dir, filename_base + '.pdf')
    pdf_abs = os.path.abspath(pdf_path)
    doc_word = word.Documents.Open(docx_abs)
    doc_word.SaveAs(pdf_abs, FileFormat=17)
    doc_word.Close()
    word.Quit()
    return new_docx_path, pdf_path

def open_borisfx_dialog(parent_root):
    dialog = tk.Toplevel(parent_root)
    dialog.title("BorisFX 인증서 만들기")
    dialog.geometry("500x650")
    dialog.configure(bg="#F7F9FB")
    dialog.resizable(False, False)
    # 카드 프레임
    card = tk.Frame(dialog, bg="white", relief="solid", bd=1)
    card.place(relx=0.5, rely=0.5, anchor="center", width=460, height=580)
    card.grid_propagate(False)
    # 회사명
    tk.Label(card, text="회사명", font=("맑은 고딕", 11), bg="white", fg="#374151", anchor='e', width=10).grid(row=0, column=0, padx=10, pady=10, sticky='e')
    entry_company = tk.Entry(card, font=("맑은 고딕", 11), bg="#F1F5F9", relief="flat", width=28)
    entry_company.grid(row=0, column=1, padx=10, pady=10, sticky='w')
    # 제품명
    tk.Label(card, text="제품명", font=("맑은 고딕", 11), bg="white", fg="#374151", anchor='e', width=10).grid(row=1, column=0, padx=10, pady=10, sticky='e')
    entry_product = tk.Entry(card, font=("맑은 고딕", 11), bg="#F1F5F9", relief="flat", width=28)
    entry_product.grid(row=1, column=1, padx=10, pady=10, sticky='w')
    # 라이선스 타입
    tk.Label(card, text="라이선스 타입", font=("맑은 고딕", 11), bg="white", fg="#374151", anchor='e', width=10).grid(row=2, column=0, padx=10, pady=10, sticky='e')
    license_type_var = tk.StringVar(value="Annual Subscription")
    radio_frame = tk.Frame(card, bg="white")
    radio_frame.grid(row=2, column=1, padx=10, pady=10, sticky='w')
    tk.Radiobutton(radio_frame, text="Annual Subscription", variable=license_type_var, value="Annual Subscription", font=("맑은 고딕", 10), bg="white").pack(side='left', padx=5)
    tk.Radiobutton(radio_frame, text="Perpetual", variable=license_type_var, value="Perpetual", font=("맑은 고딕", 10), bg="white").pack(side='left', padx=5)
    # 라이선스 정보 입력
    tk.Label(card, text="내용 붙여넣기", font=("맑은 고딕", 11), bg="white", fg="#374151", anchor='ne', width=10).grid(row=3, column=0, padx=10, pady=10, sticky='ne')
    text_content = scrolledtext.ScrolledText(card, width=36, height=5, font=("맑은 고딕", 10), bg="#F1F5F9", relief="flat")
    text_content.grid(row=3, column=1, padx=10, pady=10, sticky='w')
    # 버튼 프레임
    btn_frame = tk.Frame(card, bg="white")
    btn_frame.grid(row=4, column=0, columnspan=2, pady=18)
    def on_create():
        company = entry_company.get().strip()
        product = entry_product.get().strip()
        pasted = text_content.get("1.0", tk.END).strip()
        license_type = license_type_var.get()
        if not company or not product or not pasted:
            messagebox.showwarning("입력 오류", "모든 항목을 입력해주세요.")
            return
        try:
            docx_path, pdf_path = create_certificate(company, product, pasted, license_type)
            messagebox.showinfo("완료", f"DOCX 파일: {docx_path}\nPDF 파일: {pdf_path}\n\n[확인]을 누르면 PDF가 저장된 폴더가 열립니다.")
            try:
                folder = os.path.dirname(pdf_path)
                os.startfile(folder)
            except Exception as e:
                messagebox.showwarning("폴더 열기 실패", f"폴더를 열 수 없습니다: {e}")
            dialog.destroy()
        except Exception as e:
            messagebox.showerror("오류 발생", str(e))
    create_btn = tk.Button(btn_frame, text="만들기", font=("맑은 고딕", 11, "bold"), bg="#3B82F6", fg="white", relief="flat", activebackground="#2563EB", activeforeground="white", command=on_create, width=10, height=1)
    create_btn.pack(side='left', padx=10)
    create_btn.bind("<Enter>", lambda e: create_btn.config(bg="#2563EB"))
    create_btn.bind("<Leave>", lambda e: create_btn.config(bg="#3B82F6"))
    cancel_btn = tk.Button(btn_frame, text="취소", font=("맑은 고딕", 11, "bold"), bg="#94A3B8", fg="white", relief="flat", activebackground="#64748B", activeforeground="white", command=dialog.destroy, width=10, height=1)
    cancel_btn.pack(side='left', padx=10)
    cancel_btn.bind("<Enter>", lambda e: cancel_btn.config(bg="#64748B"))
    cancel_btn.bind("<Leave>", lambda e: cancel_btn.config(bg="#94A3B8"))
    # 사용 방법 안내 텍스트 (카드 하단)
    info_label = tk.Label(card, text="""📋 사용 방법:
1. 회사명과 제품명을 입력합니다.
2. 라이선스 정보를 텍스트 박스에 붙여넣습니다.
3. '만들기' 버튼을 클릭하여 DOCX와 PDF 파일을 생성합니다.

📁 생성된 파일은 'license_certificate/BorisFX/' 폴더에 저장됩니다.""", 
                         justify='left', anchor='nw', font=("맑은 고딕", 9), bg="white", wraplength=440)
    info_label.grid(row=5, column=0, columnspan=2, sticky='nw', padx=10, pady=(10, 10))
    dialog.transient(parent_root)
    dialog.grab_set()
    parent_root.wait_window(dialog) 