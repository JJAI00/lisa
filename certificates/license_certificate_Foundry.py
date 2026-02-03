import os
import sys
import datetime
import tkinter as tk
from tkinter import messagebox, filedialog, scrolledtext
import openpyxl
import win32com.client as wc
import shutil
import subprocess
import re
from openpyxl.cell.cell import MergedCell

def resource_path(relative_path):
    import sys, os
    if getattr(sys, 'frozen', False):
        base_path = os.path.dirname(sys.executable)
    else:
        base_path = os.path.abspath('.')
    return os.path.join(base_path, relative_path)

def resource_root():
    import sys, os
    if getattr(sys, 'frozen', False):
        return os.path.dirname(sys.executable)
    else:
        return os.path.abspath('.')

DATA_PATH = resource_path('license_certificate/Foundry/Foundry_LicenseCertificate_data.xlsx')
ORIGIN_PATH = resource_path('license_certificate/Foundry/Foundry_LicenseCertificate_origin.xlsx')
MSG_TEMPLATE_PATH = resource_path('license_certificate/Foundry/QVREX {Company Name} 라이선스 전달의 건.msg')

# --- GUI ---
def open_foundry_dialog(parent_root):
    dialog = tk.Toplevel(parent_root)
    dialog.title("Foundry 인증서/메일 자동화")
    dialog.geometry("520x600")
    dialog.configure(bg="#F7F9FB")
    dialog.resizable(False, False)

    # 카드 프레임
    card = tk.Frame(dialog, bg="white", relief="solid", bd=1)
    card.place(relx=0.5, rely=0.5, anchor="center", width=480, height=480)
    card.grid_propagate(False)

    # 데이터 붙여넣기 버튼
    def open_data_excel():
        try:
            os.startfile(DATA_PATH)
        except Exception as e:
            messagebox.showerror("엑셀 열기 오류", str(e))
    tk.Button(card, text="데이터 붙여넣기 (엑셀 열기)", command=open_data_excel, font=("맑은 고딕", 10, "bold"), bg="#3B82F6", fg="white", relief="flat", padx=10, pady=6).grid(row=0, column=0, columnspan=2, pady=(10,10), padx=10, sticky='ew')

    # 라이선스 타입 라디오 버튼
    type_var = tk.StringVar(value="Floating")
    tk.Label(card, text="라이선스 타입", font=("맑은 고딕", 11), bg="white", fg="#374151", anchor='e', width=12).grid(row=1, column=0, padx=10, pady=8, sticky='e')
    type_frame = tk.Frame(card, bg="white")
    type_frame.grid(row=1, column=1, padx=10, pady=8, sticky='w')
    for t in ["Floating", "Node-Locked", "LogIn-Based"]:
        tk.Radiobutton(type_frame, text=t, variable=type_var, value=t, font=("맑은 고딕", 10), bg="white").pack(side='left', padx=8)

    # 라이선스 정보 입력란
    tk.Label(card, text="라이선스 정보 (여러 줄 텍스트, .lic로 저장)", font=("맑은 고딕", 10), bg="white", fg="#374151").grid(row=2, column=0, columnspan=2, pady=(15,2), padx=10, sticky='w')
    license_text = tk.Text(card, font=("맑은 고딕", 10), bg="#F1F5F9", relief="flat", width=48, height=6)
    license_text.grid(row=3, column=0, columnspan=2, padx=10, pady=(0,10), sticky='ew')

    # 버튼 프레임
    btn_frame = tk.Frame(card, bg="white")
    btn_frame.grid(row=4, column=0, columnspan=2, pady=18)
    def on_cancel():
        dialog.destroy()
    def on_create():
        # 1. 데이터 엑셀 읽기
        try:
            wb = openpyxl.load_workbook(DATA_PATH)
            ws = wb.active
            if ws is None:
                raise RuntimeError('엑셀 워크시트를 찾을 수 없습니다.')
        except Exception as e:
            messagebox.showerror("엑셀 읽기 오류", str(e))
            return
        # 2. 데이터 추출 (첫 행: C=회사명, G=품목명, H=수량, U=SystemID, T=만료일, Y=Address1, Z=Address2, W=Email, X=Tel)
        company = ws['C2'].value if ws['C2'] is not None else ''
        address1 = ws['Y2'].value if ws['Y2'] is not None else ''
        address2 = ws['Z2'].value if ws['Z2'] is not None else ''
        email = ws['W2'].value if ws['W2'] is not None else ''
        contact = ws['X2'].value if ws['X2'] is not None else ''
        today_str = datetime.date.today().strftime('%Y-%m-%d')
        # 표 데이터 (여러 줄)
        # 헤더에서 각 열 인덱스 동적 추출 (정확한 위치)
        header_row = next(ws.iter_rows(min_row=1, max_row=1))
        header = [cell.value for cell in header_row]
        product_idx = header.index('품목명') if '품목명' in header else 6
        qty_idx = header.index('수량') if '수량' in header else 7
        expdate_idx = header.index('만료일') if '만료일' in header else 19
        # SystemID는 'SystemID' 또는 'System ID' 허용
        if 'SystemID' in header:
            systemid_idx = header.index('SystemID')
        elif 'System ID' in header:
            systemid_idx = header.index('System ID')
        else:
            systemid_idx = 20  # U열 (21번째 열, 인덱스 20)
        products, qtys, expdates = [], [], []
        for row in ws.iter_rows(min_row=2):
            g = row[product_idx].value if len(row) > product_idx else ''
            h = row[qty_idx].value if len(row) > qty_idx else ''
            t = row[expdate_idx].value if len(row) > expdate_idx else ''
            if g: products.append(str(g))
            if h: qtys.append(str(h))
            if t: 
                # 날짜 형식을 YYYY-MM-DD로 변환
                if isinstance(t, datetime.datetime):
                    expdates.append(t.strftime('%Y-%m-%d'))
                elif isinstance(t, datetime.date):
                    expdates.append(t.strftime('%Y-%m-%d'))
                else:
                    expdates.append(str(t))
        # SystemID만 병합 셀 처리
        systemids = []
        seen_merged = set()
        for row in ws.iter_rows(min_row=2):
            cell = row[systemid_idx]
            if isinstance(cell, MergedCell):
                continue  # 병합된 셀의 첫 셀만 사용
            merged = False
            for rng in ws.merged_cells.ranges:
                if cell.coordinate in rng:
                    if rng.start_cell.coordinate == cell.coordinate:
                        if rng.coord not in seen_merged:
                            seen_merged.add(rng.coord)
                            if cell.value:
                                systemids.append(str(cell.value))
                    merged = True
                    break
            if not merged:
                if cell.value:
                    systemids.append(str(cell.value))
        # SystemID가 병합된 셀에서 가져온 경우, Product 줄 수만큼 반복
        if len(systemids) == 1 and len(products) > 1:
            systemids = systemids * len(products)
        license_type = type_var.get()
        # 3. origin.xlsx 복사/치환
        base_dir = os.path.abspath(".")
        output_dir = os.path.join(resource_root(), 'license_certificate', 'Foundry')
        os.makedirs(output_dir, exist_ok=True)
        safe_company = company.replace(' ', '')
        filename_base = f"Foundry_LicenseCertificate_{safe_company}_{today_str}"
        new_xlsx_path = os.path.join(output_dir, filename_base + '.xlsx')
        shutil.copy(ORIGIN_PATH, new_xlsx_path)
        wb2 = openpyxl.load_workbook(new_xlsx_path)
        ws2 = wb2.active
        if ws2 is None:
            raise RuntimeError('엑셀 워크시트를 찾을 수 없습니다.')
        for row in ws2.iter_rows():
            if row is None:
                continue
            for cell in row:
                if type(cell.value) is str:
                    cell.value = cell.value.replace('{Today}', today_str)
                    cell.value = cell.value.replace('{Company Name}', company)
                    cell.value = cell.value.replace('{Address 1}', address1)
                    cell.value = cell.value.replace('{Address 2}', address2)
                    cell.value = cell.value.replace('{Email}', email)
                    cell.value = cell.value.replace('{Contact}', contact)
        # 표 데이터 입력 (A14~: Product, E14~: Qty, F14~: SystemID, G14~: Type, H14~: ExpDate)
        def is_merged_cell_first_or_not_merged(ws, cell_addr):
            for merged_range in ws.merged_cells.ranges:
                if cell_addr in merged_range:
                    return cell_addr == merged_range.start_cell.coordinate
            return True  # 병합이 아닌 셀
        for i in range(len(products)):
            a_addr = f'A{14+i}'
            if is_merged_cell_first_or_not_merged(ws2, a_addr):
                ws2[a_addr] = products[i]
            e_addr = f'E{14+i}'
            if is_merged_cell_first_or_not_merged(ws2, e_addr):
                ws2[e_addr] = qtys[i] if i < len(qtys) else ''
            f_addr = f'F{14+i}'
            if is_merged_cell_first_or_not_merged(ws2, f_addr):
                ws2[f_addr] = systemids[i] if i < len(systemids) else ''
            g_addr = f'G{14+i}'
            if is_merged_cell_first_or_not_merged(ws2, g_addr):
                ws2[g_addr] = license_type
            h_addr = f'H{14+i}'
            if is_merged_cell_first_or_not_merged(ws2, h_addr):
                ws2[h_addr] = expdates[i] if i < len(expdates) else ''
        wb2.save(new_xlsx_path)
        # PDF 변환
        excel = wc.DispatchEx('Excel.Application')
        excel.Visible = False
        wb_com = excel.Workbooks.Open(os.path.abspath(new_xlsx_path))
        pdf_path = os.path.join(output_dir, filename_base + '.pdf')
        wb_com.ExportAsFixedFormat(0, pdf_path)
        wb_com.Close(False)
        excel.Quit()
        # 4. 라이선스 정보 .lic 저장
        lic_text = license_text.get("1.0", tk.END).strip()
        lic_path = os.path.join(output_dir, f"{company}_Foundry_{today_str}.lic")
        with open(lic_path, 'w', encoding='utf-8') as f:
            f.write(lic_text)
        # 5. Outlook에서 msg 템플릿 기반 메일 생성, 치환, 첨부
        try:
            import win32com.client as win32
            outlook = win32.Dispatch("Outlook.Application")
            msg = outlook.CreateItemFromTemplate(MSG_TEMPLATE_PATH)
            # 여러 줄 데이터 본문 치환
            product_lines = '<br>'.join(products)
            qty_lines = '<br>'.join(qtys)
            systemid_lines = '<br>'.join(systemids)
            expdate_lines = '<br>'.join(expdates)
            type_lines = '<br>'.join([license_type]*len(products))
            # 라이선스 정보 치환용 텍스트
            license_key_text = license_text.get("1.0", tk.END).strip().replace('\n', '<br>')
            msg.Subject = msg.Subject.replace('{Company Name}', company)
            html_body = msg.HTMLBody
            html_body = html_body.replace('{Company Name}', company)
            html_body = html_body.replace('{Product}', product_lines)
            html_body = html_body.replace('{Qty}', qty_lines)
            html_body = html_body.replace('{Type}', type_lines)
            html_body = html_body.replace('{ExpDate}', expdate_lines)
            html_body = html_body.replace('{License Key}', license_key_text)
            msg.HTMLBody = html_body
            # 첨부파일 2종만 추가
            msg.Attachments.Add(pdf_path)
            msg.Attachments.Add(lic_path)
            msg.Display()
        except Exception as e:
            messagebox.showwarning("메일 창 열기 실패", f"Outlook 메일 창을 띄울 수 없습니다: {e}")
        messagebox.showinfo("완료", f"엑셀/PDF/라이선스가 생성되었습니다!\n\n엑셀: {new_xlsx_path}\nPDF: {pdf_path}\n라이선스: {lic_path}")
        try:
            os.startfile(output_dir)
        except Exception:
            pass
        dialog.destroy()
    def style_btn(btn, bg, hover_bg):
        btn.configure(font=("맑은 고딕", 11, "bold"), fg="white", bg=bg, relief="flat", padx=20, pady=6)
        btn.bind('<Enter>', lambda e: btn.configure(bg=hover_bg))
        btn.bind('<Leave>', lambda e: btn.configure(bg=bg))
    make_btn = tk.Button(btn_frame, text="만들기", command=on_create)
    make_btn.pack(side='left', padx=10)
    style_btn(make_btn, "#3B82F6", "#2563EB")
    cancel_btn = tk.Button(btn_frame, text="취소", command=on_cancel)
    cancel_btn.pack(side='left', padx=10)
    style_btn(cancel_btn, "#6B7280", "#374151")

    # 사용 방법 안내 (프레임 안으로 이동)
    info_label = tk.Label(card, text="""📋 사용 방법:
1. 데이터 엑셀을 열어 필요한 정보를 입력/붙여넣기 합니다.
2. 라이선스 타입과 정보를 입력 후 '만들기' 버튼을 클릭합니다.

📁 생성된 파일은 'license_certificate/'안에 각 제품별 폴더에 저장됩니다.""", 
                         justify='left', anchor='nw', font=("맑은 고딕", 9), bg="white", wraplength=440)
    info_label.grid(row=5, column=0, columnspan=2, pady=(10, 15), padx=10, sticky='w')

    dialog.transient(parent_root)
    dialog.grab_set()
    parent_root.wait_window(dialog) 